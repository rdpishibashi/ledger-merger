"""DXF-diff-manager の出力フォルダ名（"指番_モジュール_サイド"単位）で LedgerEntry を
グルーピングし、レビジョン横断の集計Excel（"{指番}_{モジュール}_{サイド}_all.xlsx"）を
生成するモジュール。Streamlit には依存しない。

DXF-diff-manager のZIPダウンロードファイル名の命名規則
（dxf_diff_results_Type{A/B/C}_{指番}_{モジュール}_{サイド}_{リビジョン}）に依存する。
2026-07-28以前に手動で "dxf_diff_results_Pair{A/B/C}_..." と命名されたフォルダ
（DXF-diff-manager がファイル名自動生成に対応する前の実データ）も引き続き解釈できる
よう、"Pair"/"Type" どちらのトークンも受け付ける。
"""

import io
import os
import re
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from utils.ledger_finder import DIFF_LIST_HEADERS

CENTER_ALIGNMENT = Alignment(horizontal="center")
_ENTITY_NUMBER_FORMAT = "#,##0"

GROUP_REVISION_PATTERN = re.compile(r'^dxf_diff_results_(?:Pair|Type)[A-Za-z]_(?P<group>.+)_(?P<revision>\d+)$')

# 差分方式（Type A/B/C。手動命名のPair{A/B/C}にも対応）を取り出すためのパターン。
# 指番/モジュール/サイドと異なり台帳ファイル名には含まれないため、フォルダ名のみから
# 判定する（2026-08、Work Master/Summaryへの列追加のため新設）。
DIFF_TYPE_PATTERN = re.compile(r'^dxf_diff_results_(?:Pair|Type)(?P<type>[A-Za-z])_')

# 指番・モジュール・サイドを個別に取り出すための厳密なパターン。DXF-diff-manager
# 自身の指番/モジュール/サイド入力フォーマット（app.py の SHIBAN_PATTERN /
# MODULE_PATTERN / SIDE_PATTERN、および両者を結合する MASTER_FILENAME_PATTERN）と
# 同一の正規表現を用いる。GROUP_REVISION_PATTERN は group を1文字列として緩く
# 取り出すのに対し、こちらは指番/モジュール/サイドそれぞれの妥当な書式を検証する。
# 末尾のリビジョン番号は2026-08以降DXF-diff-manager側のZIPファイル名から省略される
# ようになったため任意とする（例: dxf_diff_results_TypeA_ME24-1001-0_ZC00_405）。
SASHIBAN_MODULE_SIDE_PATTERN = re.compile(
    r'^dxf_diff_results_(?:Pair|Type)[A-Za-z]_'
    r'(?P<shiban>[A-Z]{2}\d{2}-\d{4}-\d)_(?P<module>[A-Z0-9]{4}|na)_(?P<side>[A-Z0-9]{3}|na)'
    r'(?:_(?P<revision>\d+))?$'
)

# 台帳 .xlsx のファイル名パターン。DXF-diff-manager 自身の台帳ファイル名規則
# （model.master_ledger.MASTER_FILENAME_PATTERN）と同一の正規表現。指番/モジュール/
# サイドの抽出は、この「ファイル名」を主として使う。フォルダ名（Diff Package名）は
# ZIPダウンロード時にユーザーが自由に編集できるテキスト欄に由来し、モジュール/サイドが
# 欠落しうる（2026-08、実データで確認: フォルダ名 "dxf_diff_results_TypeA_PE25-9601-0"
# にはモジュール/サイドが無いが、ファイル名 "PE25-9601-0_ZM00_405.xlsx" は正しく
# 持っていた）のに対し、台帳ファイル名は常にこの規則で作られるため信頼できる。
LEDGER_FILENAME_PATTERN = re.compile(
    r'^(?P<shiban>[A-Z]{2}\d{2}-\d{4}-\d)_(?P<module>[A-Z0-9]{4}|na)_(?P<side>[A-Z0-9]{3}|na)'
    r'(?:[_-].*)?\.xlsx$'
)

_CHILD_COL = DIFF_LIST_HEADERS.index("Child")
_RECORDED_DATE_COL = DIFF_LIST_HEADERS.index("Recorded Date")
_ENTITY_COLS = ("Deleted Entities", "Added Entities", "Diff Entities", "Unchanged Entities", "Total Entities")

# Summaryシートの行構成（セクション見出し, 項目ラベル）。項目ラベルは DXF-diff-manager
# 自身のSummaryシートの表記（Type A: all_in_one）をそのまま使う。Type B/C
# （流用先図面総数/流用先図面 図形総数）の文言には対応しない（既知の制約）。
# 「完全新規図面数」「新規作成率 [%]」は2026-08追加（DXF-diff-manager Summaryシートの
# 対応する2指標と同じ相対位置：差分抽出ペア数の直下・流用率[%]の直下）。
SUMMARY_ROWS = (
    ("エンティティ統計", "削除図形 総数"),
    (None, "追加図形 総数"),
    (None, "変更（追加+削除）図形 総数"),
    (None, "変更なし図形 総数"),
    (None, "アップロード図面 図形総数"),
    (None, "図形変更率 [%]"),
    ("図面統計", "アップロード図面総数"),
    (None, "差分抽出ペア数"),
    (None, "完全新規図面数"),
    (None, "流用率 [%]"),
    (None, "新規作成率 [%]"),
)

# Summaryシートの項目ラベル → LedgerEntry.summary_values のキー（utils.ledger_finder の
# SUMMARY_LABELS/OPTIONAL_SUMMARY_LABELS、Ledger-merger自身の統合Excel列名）への対応。
_CANONICAL_BY_DISPLAY_LABEL = {
    "削除図形 総数": "削除図形数 合計",
    "追加図形 総数": "追加図形数 合計",
    "変更（追加+削除）図形 総数": "差分図形数 合計",
    "変更なし図形 総数": "変更なし図形数 合計",
    "アップロード図面 図形総数": "総図形数 合計",
    "図形変更率 [%]": "図形変更率 [%]",
    "アップロード図面総数": "入力図面総数",
    "差分抽出ペア数": "差分抽出ペア数",
    "完全新規図面数": "完全新規図面数",
    "流用率 [%]": "流用率 [%]",
    "新規作成率 [%]": "新規作成率 [%]",
}

_COUNT_LABELS = (
    "削除図形 総数", "追加図形 総数", "変更（追加+削除）図形 総数", "変更なし図形 総数",
    "アップロード図面 図形総数", "アップロード図面総数", "差分抽出ペア数", "完全新規図面数",
)
_PERCENT_LABELS = {"図形変更率 [%]", "流用率 [%]", "新規作成率 [%]"}


def parse_sashiban_module_side(package_name, filename=None):
    """指番・モジュール・サイドを取り出す。DXF-diff-manager 側で未入力（"na"）だった
    場合はそのまま文字列 "na" を返す。

    台帳ファイル名（filename。"{指番}_{モジュール}_{サイド}[_-suffix].xlsx"、
    LEDGER_FILENAME_PATTERN）を主として使う。フォルダ名（package_name。ZIP
    ダウンロード時にユーザーが自由編集できるテキスト欄に由来し、モジュール/サイドが
    欠落しうる）より信頼できるため（2026-08 ユーザー指摘）。

    filename が未指定、またはこの規則に一致しない場合のみ、フォルダ名の厳密パターン
    （SASHIBAN_MODULE_SIDE_PATTERN）へフォールバックする。どちらにも一致しない場合は
    (None, None, None) を返す（この機能の対象外として扱う）。
    """
    if filename:
        match = LEDGER_FILENAME_PATTERN.match(os.path.basename(filename))
        if match:
            return match['shiban'], match['module'], match['side']

    match = SASHIBAN_MODULE_SIDE_PATTERN.match(package_name)
    if not match:
        return None, None, None
    return match['shiban'], match['module'], match['side']


def parse_diff_type(package_name):
    """Diff Package（DXF-diff-manager 出力フォルダ名）から差分方式（Type A/B/C。
    手動命名の Pair{A/B/C} にも対応）を取り出す。指番/モジュール/サイドと異なり
    台帳ファイル名には含まれないため、フォルダ名のみから判定する。一致しない場合は
    None を返す（Work Master/Summary の Diff Type / 差分方式列が空欄になる）。
    """
    match = DIFF_TYPE_PATTERN.match(package_name)
    return match['type'] if match else None


def parse_group_and_revision(package_name, filename=None):
    """グループキー（指番_モジュール_サイド）とレビジョン番号（省略時は None）を
    取り出す。一致しない場合は None を返す（この機能の対象外として扱う）。

    指番/モジュール/サイドは parse_sashiban_module_side() と同じ優先順位
    （ファイル名を主、フォルダ名の厳密パターンを従）で決定する。レビジョン番号は
    フォルダ名（DXF-diff-managerのZIPダウンロードファイル名の命名規則
    dxf_diff_results_Type{A/B/C}_{指番}_{モジュール}_{サイド}_{リビジョン}、または
    リビジョン省略形〈2026-08以降〉）からのみ取得する——台帳ファイル名の末尾サフィックス
    はリビジョン以外の自由文字列でありうる（LEDGER_FILENAME_PATTERN の "_-suffix"
    部分）ため、レビジョンとして解釈しない（2026-08 ユーザー指摘）。

    指番/モジュール/サイドがどちらの規則にも一致しない場合のみ、フォルダ名の緩い
    パターン（GROUP_REVISION_PATTERN。groupを1文字列として緩く取り出す）へ
    フォールバックする。これにより、指番書式に一致しない過去の手動命名フォルダ
    （例: "..._OK_01"）の解釈は変えない。
    """
    sashiban, module, side = parse_sashiban_module_side(package_name, filename)
    if sashiban is not None:
        group = f"{sashiban}_{module}_{side}"
        strict_match = SASHIBAN_MODULE_SIDE_PATTERN.match(package_name)
        if strict_match:
            return group, strict_match['revision']
        loose_match = GROUP_REVISION_PATTERN.match(package_name)
        if loose_match:
            return group, loose_match['revision']
        return group, None

    match = GROUP_REVISION_PATTERN.match(package_name)
    if not match:
        return None
    return match['group'], match['revision']


def find_entries_with_unresolved_sashiban(entries):
    """有効な台帳として検出されたが、ファイル名・フォルダ名のどちらからも指番・
    モジュール・サイドを特定できなかった LedgerEntry を返す。

    これらのエントリは Work Master・Summary・指番_モジュール_サイド別集計から
    黙って除外される（図形変更量詳細.xlsx には Sashiban/Module/Side が空欄のまま
    引き続き含まれる）。台帳ファイル名が命名規則からずれている場合（2026-08
    ユーザー報告の実例: "NE24-0062-0_ZM00_405l.xlsx" — 末尾に区切り文字
    "_"/"-" の無い "l" が付いておりミスタイプだった）を検出し、ユーザーに知らせる
    ための一覧。
    """
    return [
        entry for entry in entries
        if parse_sashiban_module_side(entry.package_name, entry.source_path)[0] is None
    ]


def _max_recorded_date(entry):
    dates = [row[_RECORDED_DATE_COL] for row in entry.diff_list_rows if row[_RECORDED_DATE_COL] is not None]
    return max(dates) if dates else datetime.min


def group_entries(entries):
    """LedgerEntry のリストを、指番_モジュール_サイど ごと・レビジョンごとにまとめる。

    同一フォルダ（同一 package_name）に複数の有効な台帳が見つかった場合は、Diff List
    内の最大 Recorded Date が最も新しいものだけを採用する（差分抽出のやり直しで古い
    実行結果がフォルダに残っていた場合の取り違え防止。2026-07-28 の実データで実際に
    確認したケース: 図番抽出に失敗した古い実行結果と、成功した新しい実行結果が
    同じフォルダに混在していた）。**この「同一フォルダ内の複数候補から1つを選ぶ」
    処理は、指番/モジュール/サイドの算出（ファイル名を主とする parse_group_and_revision
    参照）より先に行う** — 同じフォルダ内の候補ファイルは、たとえファイル名の命名が
    異なっていても（例: 失敗した古い実行結果の "..._na_na.xlsx" と、成功した新しい
    実行結果の "..._ZC00_405.xlsx"）同一の実行対象として競合させ、勝者を決めてから
    その勝者のファイル名で指番/モジュール/サイドを決定する（2026-08、フォルダ名を
    ファイル名優先に変更した際、先に指番/モジュール/サイドで競合グループを分けて
    しまうと、この「同一フォルダ内の取り違え防止」が機能しなくなる回帰を作ったため
    修正）。

    Returns:
        dict[group_key, list[(revision, LedgerEntry)]]（レビジョン文字列の昇順。
        リビジョン省略形〈revision が None〉のエントリは末尾に並べる）
    """
    by_folder = defaultdict(list)
    for entry in entries:
        by_folder[entry.package_name].append(entry)

    groups = defaultdict(list)
    for _package_name, folder_entries in by_folder.items():
        latest_entry = max(folder_entries, key=_max_recorded_date)
        parsed = parse_group_and_revision(latest_entry.package_name, latest_entry.source_path)
        if parsed is None:
            continue
        group_key, revision = parsed
        groups[group_key].append((revision, latest_entry))

    for group_key in groups:
        groups[group_key].sort(key=lambda item: (item[0] is None, item[0] or ""))

    return dict(groups)


def aggregate_diff_list_by_child(revision_entries):
    """revision_entries（[(revision, LedgerEntry), ...]）の全 diff_list_rows を
    Child ごとに集計する。

    Deleted/Added/Diff/Unchanged/Total Entities の5列は全レビジョンにわたって
    単純合計する（'n/a' の値は数値でないためスキップし、ある Child の全出現が
    'n/a' の列はそのまま 'n/a' として残す）。非数値列（Parent/Relation/Title/
    Subtitle/Recorded Date/Note）は、Recorded Date が最も新しい行の値を採用する。

    Returns:
        list[tuple]（DIFF_LIST_HEADERS 12列、Childの昇順）
    """
    rows_by_child = defaultdict(list)
    for _revision, entry in revision_entries:
        for row in entry.diff_list_rows:
            rows_by_child[row[_CHILD_COL]].append(row)

    aggregated = []
    for child in sorted(rows_by_child.keys()):
        rows = rows_by_child[child]
        latest_row = max(rows, key=lambda r: r[_RECORDED_DATE_COL] or datetime.min)

        out_row = list(latest_row)
        for col_name in _ENTITY_COLS:
            col_idx = DIFF_LIST_HEADERS.index(col_name)
            numeric_values = [r[col_idx] for r in rows if isinstance(r[col_idx], (int, float))]
            out_row[col_idx] = sum(numeric_values) if numeric_values else 'n/a'
        aggregated.append(tuple(out_row))

    return aggregated


def _revision_value(entry, display_label):
    """entry.summary_values から表示ラベルに対応する値を取得する。「完全新規図面数」
    「新規作成率 [%]」は任意ラベル（utils.ledger_finder.OPTIONAL_SUMMARY_LABELS）の
    ため、これらを持たない旧バージョンの台帳では値が存在しない。その場合は0として
    扱う（Summaryシートの合計・レビジョン列は常に数値で表示するため）。
    """
    return entry.summary_values.get(_CANONICAL_BY_DISPLAY_LABEL[display_label], 0)


def _total_value(revision_entries, display_label):
    if display_label == "図形変更率 [%]":
        total_diff = _total_value(revision_entries, "変更（追加+削除）図形 総数")
        total_entities = _total_value(revision_entries, "アップロード図面 図形総数")
        return (total_diff / total_entities) if total_entities else 0.0
    if display_label == "流用率 [%]":
        total_pairs = _total_value(revision_entries, "差分抽出ペア数")
        total_drawings = _total_value(revision_entries, "アップロード図面総数")
        return (total_pairs / total_drawings) if total_drawings else 0.0
    if display_label == "新規作成率 [%]":
        total_brand_new = _total_value(revision_entries, "完全新規図面数")
        total_drawings = _total_value(revision_entries, "アップロード図面総数")
        return (total_brand_new / total_drawings) if total_drawings else 0.0
    return sum(_revision_value(entry, display_label) for _revision, entry in revision_entries)


def aggregate_input_drawing_totals_by_sashiban_and_diff_type(entries):
    """entries（今回のZIP入力から得たLedgerEntryのリスト）を指番_モジュール_サイド
    単位でグルーピングし、各グループの「アップロード図面総数」TOTAL値
    （指番_モジュール_サイド別集計フォルダの "{group_key}_all.xlsx" Summaryシートの
    「アップロード図面総数」行・TOTAL列と同値。"アップロード図面総数" はカウント系
    ラベルのため _total_value() の実体は単純合計だが、ここでは非数値・欠損値
    （summary_values にキーが無い場合等）を安全に0として扱うため独自に集計する）を、
    (指番, 差分方式) ごとに合算する（2026-08、Summaryの差分方式列追加に伴い
    指番単独のキーから拡張。差分方式が同一指番内で異なる場合は別グループとして
    分ける——ユーザー確認済み仕様）。命名規則に一致しないグループは対象外。

    Returns:
        dict[(sashiban, diff_type), int|float]
    """
    groups = group_entries(entries)
    totals = defaultdict(int)
    for group_key, revision_entries in groups.items():
        representative_entry = revision_entries[0][1]
        sashiban, _module, _side = parse_sashiban_module_side(
            representative_entry.package_name, representative_entry.source_path)
        if sashiban is None:
            continue
        diff_type = parse_diff_type(representative_entry.package_name)
        for _revision, entry in revision_entries:
            value = entry.summary_values.get(_CANONICAL_BY_DISPLAY_LABEL["アップロード図面総数"])
            if isinstance(value, (int, float)):
                totals[(sashiban, diff_type)] += value
    return dict(totals)


def build_group_workbook(group_key, revision_entries):
    """1つのグループ（指番_モジュール_サイド）について、レビジョン横断の
    Summary + Diff List を持つ統合Excel（bytes）を生成する。"""
    wb = Workbook()

    # --- Summary シート（先に追加してタブ順を先頭にする） ---
    summary_ws = wb.active
    summary_ws.title = "Summary"

    # リビジョン省略形（revision が全件 None）のグループは、レビジョン列を出さず
    # TOTAL列のみとする。一部だけ省略形が混在する場合（通常は発生しないが、命名
    # 規則の移行期に起こりうる）は、見出しを "-" として列自体は残す。
    revisions = [revision for revision, _entry in revision_entries]
    show_revision_columns = any(r is not None for r in revisions)
    revision_col_count = len(revisions) if show_revision_columns else 0

    header_row = [None, None, "TOTAL"]
    if show_revision_columns:
        header_row += [f'"{r if r is not None else "-"}"' for r in revisions]
    summary_ws.append(header_row)
    for cell in summary_ws[1]:
        cell.font = Font(bold=True)

    for section, label in SUMMARY_ROWS:
        total = _total_value(revision_entries, label)
        row = [section, label, total]
        if show_revision_columns:
            row += [_revision_value(entry, label) for _revision, entry in revision_entries]
        summary_ws.append(row)

        row_idx = summary_ws.max_row
        number_format = "0.00%" if label in _PERCENT_LABELS else "#,##0"
        for col_idx in range(3, 3 + 1 + revision_col_count):
            summary_ws.cell(row=row_idx, column=col_idx).number_format = number_format

    summary_ws.column_dimensions["A"].width = 18
    summary_ws.column_dimensions["B"].width = 24
    for col_idx in range(3, 3 + 1 + revision_col_count):
        summary_ws.column_dimensions[summary_ws.cell(row=1, column=col_idx).column_letter].width = 12

    # --- Diff List シート（"Diff Package"・Summary9項目列は含めない。Childごとに集計済み） ---
    diff_ws = wb.create_sheet("Diff List")
    diff_ws.append(DIFF_LIST_HEADERS)
    for cell in diff_ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = CENTER_ALIGNMENT
    diff_ws.freeze_panes = "A2"

    recorded_date_col = DIFF_LIST_HEADERS.index("Recorded Date") + 1
    entity_cols = [DIFF_LIST_HEADERS.index(label) + 1 for label in _ENTITY_COLS]
    for row in aggregate_diff_list_by_child(revision_entries):
        diff_ws.append(row)
        row_idx = diff_ws.max_row
        diff_ws.cell(row=row_idx, column=recorded_date_col).number_format = "YYYY-MM-DD HH:MM:SS"
        # "* Entities" 列（数値/'n/a' 混在）はカンマ区切り＋中央揃いで表示位置を揃える
        for col in entity_cols:
            cell = diff_ws.cell(row=row_idx, column=col)
            cell.number_format = _ENTITY_NUMBER_FORMAT
            cell.alignment = CENTER_ALIGNMENT

    for col_idx, header in enumerate(DIFF_LIST_HEADERS, start=1):
        width = max(len(str(header)) + 2, 12)
        diff_ws.column_dimensions[diff_ws.cell(row=1, column=col_idx).column_letter].width = width

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def build_group_workbooks(entries):
    """entries から検出できる全グループについて、"{指番}_{モジュール}_{サイド}_all.xlsx"
    を生成する。

    Returns:
        dict[filename, bytes]（グループキーの昇順）
    """
    groups = group_entries(entries)
    return {
        f"{group_key}_all.xlsx": build_group_workbook(group_key, groups[group_key])
        for group_key in sorted(groups)
    }
