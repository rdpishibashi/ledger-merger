"""複数の台帳ファイル（LedgerEntry）の Diff List を1つの Excel に統合するモジュール。
Streamlit には依存しない。"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from utils.group_summary_builder import parse_diff_type, parse_sashiban_module_side
from utils.ledger_finder import DIFF_LIST_HEADERS

FIRST_ROW_FONT = Font(color="FF000000")
OTHER_ROW_FONT = Font(color="FFA6A6A6")

CENTER_ALIGNMENT = Alignment(horizontal="center")

ENTITY_LABELS = ("Deleted Entities", "Added Entities", "Diff Entities", "Unchanged Entities", "Total Entities")

# LedgerEntry.summary_values の正規キー（utils.ledger_finder.SUMMARY_LABELS の一部）→
# 図形変更量詳細.xlsx に出力する表示列名。エンティティ統計5項目のみに絞り、図面統計系
# （入力図面総数・差分抽出ペア数・流用率 [%]）は含めない（2026-08、ユーザー要望により
# 図面統計は指番_モジュール_サイド別集計・統合図面管理台帳側の指標に一本化し、こちらの
# 変更量詳細は図形の変更内容に絞った。併せて「差分図形数 合計」→「変更図形数 合計」、
# 「総図形数 合計」→「図形総数 合計」に表示名を変更）。
_MERGED_SUMMARY_COLUMNS = (
    ("削除図形数 合計", "削除図形数 合計"),
    ("追加図形数 合計", "追加図形数 合計"),
    ("差分図形数 合計", "変更図形数 合計"),
    ("総図形数 合計", "図形総数 合計"),
    ("図形変更率 [%]", "図形変更率 [%]"),
)
_MERGED_SUMMARY_CANONICAL_KEYS = tuple(canonical for canonical, _display in _MERGED_SUMMARY_COLUMNS)
_MERGED_SUMMARY_DISPLAY_LABELS = tuple(display for _canonical, display in _MERGED_SUMMARY_COLUMNS)

PERCENT_LABELS = {"図形変更率 [%]"}

# "Sashiban"/"Module"/"Side" は台帳ファイル名を主・Diff Package（DXF-diff-manager
# 出力フォルダ名）を従として parse_sashiban_module_side() で逆算した指番/モジュール/
# サイド。"Diff Package" 自体は
# 参照情報として最終列に残す（2026-07-31、ユーザー要望によりChildの前に指番系3列を
# 追加し、Diff Packageを先頭から最終列へ移動）。
# 明示的な列順のタプルとして定義（2026-08、DIFF_LIST_HEADERSの並びをそのまま転記する
# 方式から変更。Subtitleと Deleted Entitiesの間に Diff Type を追加し、Note・
# Recorded Date を 図形変更率 [%] の後ろ・Diff Packageの前へ移動——Master/Work Master
# と同じ並び替えパターン）。
OUTPUT_HEADERS = (
    "Sashiban", "Module", "Side", "Child", "Parent", "Relation", "Title", "Subtitle", "Diff Type",
) + ENTITY_LABELS + _MERGED_SUMMARY_DISPLAY_LABELS + ("Note", "Recorded Date", "Diff Package")
RECORDED_DATE_COL = OUTPUT_HEADERS.index("Recorded Date") + 1
ENTITY_COLS = [OUTPUT_HEADERS.index(label) + 1 for label in ENTITY_LABELS]
DIFF_PACKAGE_COL = OUTPUT_HEADERS.index("Diff Package") + 1
DIFF_TYPE_COL = OUTPUT_HEADERS.index("Diff Type") + 1
_SUMMARY_COLS = [OUTPUT_HEADERS.index(label) + 1 for label in _MERGED_SUMMARY_DISPLAY_LABELS]
_CHILD_COL = DIFF_LIST_HEADERS.index("Child")
_PARENT_COL = DIFF_LIST_HEADERS.index("Parent")
_RELATION_COL = DIFF_LIST_HEADERS.index("Relation")
_TITLE_COL = DIFF_LIST_HEADERS.index("Title")
_SUBTITLE_COL = DIFF_LIST_HEADERS.index("Subtitle")
_SRC_RECORDED_DATE_COL = DIFF_LIST_HEADERS.index("Recorded Date")
_NOTE_COL = DIFF_LIST_HEADERS.index("Note")
_DELETED_COL = DIFF_LIST_HEADERS.index("Deleted Entities")
_ADDED_COL = DIFF_LIST_HEADERS.index("Added Entities")
_DIFF_COL = DIFF_LIST_HEADERS.index("Diff Entities")
_UNCHANGED_COL = DIFF_LIST_HEADERS.index("Unchanged Entities")
_TOTAL_COL = DIFF_LIST_HEADERS.index("Total Entities")


def _sort_part(value):
    """None を安全に最後尾へ回すためのソートキー要素（Noneどうし・文字列どうしの
    比較のみになり、None と str の比較で TypeError が出ないようにする）。"""
    return (value is None, value or "")


def build_merged_workbook(entries):
    """LedgerEntry のリストから統合済み Excel を生成し bytes で返す。行は
    Sashiban → Diff Package → Module → Side → Child の昇順にソートする
    （2026-08、ユーザー要望）。Sashiban/Module/Side を逆算できないエントリ
    （`parse_sashiban_module_side()` が None を返す場合）は末尾に回る。

    各 LedgerEntry の行はソート後も (Sashiban, Diff Package, Module, Side) が
    同一のため連続してまとまる。「ブロック内最初の行のみ集計値・Diff Package列を
    黒字にする」という既存の表示規則は、ソート後にその LedgerEntry の行として
    最初に現れる行（＝ソート後にChildが最小の行）に適用される。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Diff List"

    ws.append(OUTPUT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = CENTER_ALIGNMENT
    ws.freeze_panes = "A2"

    flat_rows = []
    for entry in entries:
        sashiban, module, side = parse_sashiban_module_side(entry.package_name, entry.source_path)
        diff_type = parse_diff_type(entry.package_name)
        for diff_row in entry.diff_list_rows:
            sort_key = (
                _sort_part(sashiban), entry.package_name,
                _sort_part(module), _sort_part(side), diff_row[_CHILD_COL],
            )
            flat_rows.append((sort_key, entry, sashiban, module, side, diff_type, diff_row))
    flat_rows.sort(key=lambda item: item[0])

    row_idx = 2
    seen_entry_ids = set()
    for _sort_key, entry, sashiban, module, side, diff_type, diff_row in flat_rows:
        is_first_in_block = id(entry) not in seen_entry_ids
        seen_entry_ids.add(id(entry))

        row = [
            sashiban, module, side, diff_row[_CHILD_COL], diff_row[_PARENT_COL], diff_row[_RELATION_COL],
            diff_row[_TITLE_COL], diff_row[_SUBTITLE_COL], diff_type,
            diff_row[_DELETED_COL], diff_row[_ADDED_COL], diff_row[_DIFF_COL],
            diff_row[_UNCHANGED_COL], diff_row[_TOTAL_COL],
        ]
        if is_first_in_block:
            row += [entry.summary_values[canonical] for canonical in _MERGED_SUMMARY_CANONICAL_KEYS]
        else:
            row += [None] * len(_MERGED_SUMMARY_CANONICAL_KEYS)
        row += [diff_row[_NOTE_COL], diff_row[_SRC_RECORDED_DATE_COL], entry.package_name]
        ws.append(row)

        package_cell = ws.cell(row=row_idx, column=DIFF_PACKAGE_COL)
        package_cell.font = FIRST_ROW_FONT if is_first_in_block else OTHER_ROW_FONT

        ws.cell(row=row_idx, column=RECORDED_DATE_COL).number_format = "YYYY-MM-DD HH:MM:SS"
        ws.cell(row=row_idx, column=DIFF_TYPE_COL).alignment = CENTER_ALIGNMENT

        # "* Entities" 列（数値/'n/a' 混在）はカンマ区切り＋中央揃いで表示位置を揃える
        for col in ENTITY_COLS:
            cell = ws.cell(row=row_idx, column=col)
            cell.number_format = "#,##0"
            cell.alignment = CENTER_ALIGNMENT

        if is_first_in_block:
            for col, label in zip(_SUMMARY_COLS, _MERGED_SUMMARY_DISPLAY_LABELS):
                cell = ws.cell(row=row_idx, column=col)
                cell.number_format = "0.00%" if label in PERCENT_LABELS else "#,##0"
                cell.alignment = CENTER_ALIGNMENT

        row_idx += 1

    for col_idx, header in enumerate(OUTPUT_HEADERS, start=1):
        width = max(len(str(header)) + 2, 12)
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
