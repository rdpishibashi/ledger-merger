"""
Bugfix regression: 2026-09-18、DXF-diff-manager が Master シートに
"Unchanged Offset Entities" 列（オフセット補正で一致した図形数）を新設したことに
伴い、Ledger-merger 側の DIFF_LIST_HEADERS が12列→13列に変わった。

不具合になり得た事象（実際に発生する前に対策）: _find_diff_list_rows() が
新形式（13列）の DIFF_LIST_HEADERS との完全一致でのみデータシートを判定すると、
2026-09-18より前に生成された旧形式（12列）の台帳ファイルは、ヘッダーが一致せず
「有効な台帳が見つからない」と黙って判定され、Ledger-merger の統合対象から
静かに抜け落ちてしまう——2026-08-30 のシート名変更（"Diff List"→"Master"）で
実際に発生しかけた問題と同じ失敗モード（tests/regression/bugfix/
test_sheet_name_agnostic_detection.py 参照）。

修正: _find_diff_list_rows() が新形式（DIFF_LIST_HEADERS）・旧形式
（LEGACY_DIFF_LIST_HEADERS）の両方のヘッダーを検出し、旧形式の場合は各データ行の
末尾に 0 を補って新形式と同じ13列に揃えてから返す（旧形式にはオフセット補正機能
自体が無かったため、実際の値として0が正しい）。

保証したいこと:
- 旧形式（12列）の台帳ファイルも find_ledger_files() で検出されること
  （folders_without_ledger に落ちない）
- 検出された diff_list_rows は新形式と同じ13列になり、末尾（Unchanged Offset
  Entities）が 0 であること
- 新形式（13列、実際のオフセット値あり）の台帳も引き続き正しく検出され、
  値がそのまま保たれること

実行:
    cd Ledger-merger
    python -m pytest tests/regression/bugfix/test_unchanged_offset_column_legacy_compat.py -v
"""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", ".."))

from openpyxl import Workbook

from utils.ledger_finder import DIFF_LIST_HEADERS, find_ledger_files

_LEGACY_HEADERS = [
    "Child", "Parent", "Relation", "Title", "Subtitle", "Recorded Date", "Note",
    "Deleted Entities", "Added Entities", "Diff Entities", "Unchanged Entities",
    "Total Entities",
]

_SUMMARY_ROWS = [
    ("削除図形 総数", 1), ("追加図形 総数", 2), ("変更（追加+削除）図形 総数", 3),
    ("変更なし図形 総数", 4), ("アップロード図面 図形総数", 5),
    ("図形変更率 [%]", 0.6),
]


def _build_legacy_ledger_file(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    data_ws = wb.active
    data_ws.title = "Master"
    data_ws.append(_LEGACY_HEADERS)
    data_ws.append(["C1", "P1", "RevUp", "T", "S", None, None, 1, 2, 3, 4, 10])

    summary_ws = wb.create_sheet("Summary")
    for row in _SUMMARY_ROWS:
        summary_ws.append(row)

    wb.save(path)


def _build_new_format_ledger_file(path, offset_value=7):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    data_ws = wb.active
    data_ws.title = "Master"
    data_ws.append(list(DIFF_LIST_HEADERS))
    data_ws.append(["C1", "P1", "RevUp", "T", "S", None, None, 1, 2, 3, 4, 10, offset_value])

    summary_ws = wb.create_sheet("Summary")
    for row in _SUMMARY_ROWS:
        summary_ws.append(row)
    summary_ws.append(("変更なし（オフセット一致）図形 総数", offset_value))

    wb.save(path)


def test_legacy_format_ledger_still_detected_and_padded_with_zero(tmp_path):
    """旧形式（12列、Unchanged Offset Entities列なし）の台帳が黙って除外されず
    検出され、新形式と同じ13列（末尾0）に揃えられる。"""
    folder = tmp_path / "PACKAGE-LEGACY"
    _build_legacy_ledger_file(folder / "legacy_master.xlsx")

    entries, missing = find_ledger_files(str(tmp_path))

    assert missing == [], f"旧形式の台帳が黙って除外された: {missing}"
    assert len(entries) == 1
    row = entries[0].diff_list_rows[0]
    assert len(row) == len(DIFF_LIST_HEADERS) == 13
    assert row[-1] == 0  # Unchanged Offset Entities（旧形式は常に0）


def test_new_format_ledger_detected_with_offset_value_preserved(tmp_path):
    """新形式（13列）の台帳も引き続き検出され、Unchanged Offset Entities の
    実際の値がそのまま保たれる。"""
    folder = tmp_path / "PACKAGE-NEW"
    _build_new_format_ledger_file(folder / "new_master.xlsx", offset_value=7)

    entries, missing = find_ledger_files(str(tmp_path))

    assert missing == []
    assert len(entries) == 1
    row = entries[0].diff_list_rows[0]
    assert len(row) == 13
    assert row[-1] == 7


def test_legacy_and_new_format_ledgers_coexist_in_same_scan(tmp_path):
    """同一の走査対象ディレクトリに旧形式・新形式の台帳が混在していても、
    両方とも正しく検出される（DXF-diff-managerのバージョンが異なる複数の
    出力フォルダが混在する運用を想定）。"""
    _build_legacy_ledger_file(tmp_path / "PACKAGE-LEGACY" / "legacy_master.xlsx")
    _build_new_format_ledger_file(tmp_path / "PACKAGE-NEW" / "new_master.xlsx", offset_value=3)

    entries, missing = find_ledger_files(str(tmp_path))

    assert missing == []
    assert len(entries) == 2
    offsets_by_package = {e.package_name: e.diff_list_rows[0][-1] for e in entries}
    assert offsets_by_package == {"PACKAGE-LEGACY": 0, "PACKAGE-NEW": 3}


if __name__ == '__main__':
    import pytest
    sys.exit(pytest.main([__file__, "-v"]))
