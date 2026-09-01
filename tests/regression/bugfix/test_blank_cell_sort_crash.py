"""
Bugfix regression: 2026-09, アップロードされた前回の統合図面管理台帳.xlsxに
Sashiban/Module/Side/Child のいずれかが空欄のセルを持つ行が混在していると、
Master・Work Masterシート生成時のソートで
`TypeError: '<' not supported between instances of 'str' and 'NoneType'`
が発生し、アプリがクラッシュしていた（ユーザー実機で発生・報告）。

不具合: build_master_workbook() 内の _write_ledger_sheet() 呼び出しに渡す
sort_key が、キーの各要素をそのまま比較していた。今回分のエントリは
parse_sashiban_module_side() が返す値が必ず文字列（未入力時も "na"）で
None にはならないが、アップロードされた前回台帳（read_master_rows()/
read_work_master_rows() が openpyxl でセルの生値をそのまま読む）に
空欄セルがあると、そのキー要素が None になる。今回分の文字列キーと
前回分の None キーが混在した状態で sorted() が比較を試みると、
Python 3 では str と None の大小比較ができずクラッシュする。

修正: sort_key の各キー要素を `or ''` で防御し、None を空文字列として
扱うようにした（Master・Work Master 両方）。

保証したいこと:
- Master・Work Master いずれも、前回分に None を含むキーが混在していても
  クラッシュせず生成できること。
- None は空文字列として扱われ、他の値との比較で例外にならないこと
  （並び順自体の厳密な妥当性は問わない——空欄セルは元々異常データであり、
  「クラッシュしない」ことが本テストの主眼）。
"""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", ".."))

import openpyxl

from utils.master_ledger_builder import build_master_workbook, MASTER_SHEET_NAME, WORK_MASTER_SHEET_NAME


def test_work_master_sort_survives_none_in_previous_rows():
    """previous_work_master_rows に None を含むキーが混在してもクラッシュしない。"""
    previous_work_master_rows = {
        ("AA11-1111-1", None, "405", "C1", "P1"):
            ("AA11-1111-1", None, "405", "C1", "P1", "T", "S", "A", 1, 2, 3, 4, 10, None, None),
        ("AA11-1111-1", "ZM00", "405", "C2", "P2"):
            ("AA11-1111-1", "ZM00", "405", "C2", "P2", "T", "S", "A", 1, 2, 3, 4, 10, None, None),
    }

    data = build_master_workbook([], previous_work_master_rows=previous_work_master_rows)

    wb = openpyxl.load_workbook(__import__("io").BytesIO(data))
    ws = wb[WORK_MASTER_SHEET_NAME]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 3  # ヘッダー + 2行
    children = {row[3] for row in rows[1:]}
    assert children == {"C1", "C2"}


def test_master_sort_survives_none_in_previous_rows():
    """previous_master_rows に Child が None のキーが混在してもクラッシュしない。"""
    previous_master_rows = {
        (None, "P1"): (None, "P1", "流用", "T", "S", "A", 1, 2, 3, 4, 10, None, None),
        ("C2", "P2"): ("C2", "P2", "流用", "T", "S", "A", 1, 2, 3, 4, 10, None, None),
    }

    data = build_master_workbook([], previous_master_rows=previous_master_rows)

    wb = openpyxl.load_workbook(__import__("io").BytesIO(data))
    ws = wb[MASTER_SHEET_NAME]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 3  # ヘッダー + 2行


def test_both_sheets_survive_combined_none_and_string_keys():
    """MasterとWork Master双方に前回分のNoneキーが混在した状態を同時に処理してもクラッシュしない。"""
    previous_master_rows = {
        (None, None): (None, None, None, None, None, None, None, None, None, None, None, None, None),
    }
    previous_work_master_rows = {
        (None, None, None, None, None):
            (None, None, None, None, None, None, None, None, None, None, None, None, None, None, None),
    }

    data = build_master_workbook(
        [], previous_master_rows=previous_master_rows,
        previous_work_master_rows=previous_work_master_rows,
    )
    assert isinstance(data, bytes) and len(data) > 0
