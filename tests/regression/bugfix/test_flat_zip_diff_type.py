"""
Bugfix regression: 2026-10-18、フォルダを含まず台帳が直下にある ZIP
（dxf_diff_results_TypeC_YO26-0815-2_ZM99_339.zip 等、DXF-diff-manager の
ダウンロード形式）を統合すると、Master / Work Master の Diff Type と Summary の
差分方式が空欄になった。

原因: Diff Type は Diff Package 名（台帳のあるフォルダ名）の Type{A/B/C} から
判定するが、フォルダの無い ZIP では展開先の一時フォルダ名（ランダム）が
Diff Package 名になっていた。

修正: 展開先を ZIP のファイル名のフォルダにする（zip_extract_subdir）。

保証したいこと:
- フォルダ無し ZIP でも Diff Package 名が ZIP のファイル名になり Diff Type が取れる
- フォルダ有り ZIP は従来どおり内側のフォルダ名が Diff Package 名になる
- Master / Work Master / Summary の全てに Diff Type が反映される

実行:
    cd Ledger-merger
    python -m pytest tests/regression/bugfix/test_flat_zip_diff_type.py -v
"""
import io
import os
import sys
import tempfile
import zipfile

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", ".."))

import openpyxl
from openpyxl import Workbook

from utils.ledger_finder import DIFF_LIST_HEADERS, find_ledger_files, zip_extract_subdir
from utils.master_ledger_builder import build_master_workbook

ZIP_NAME = "dxf_diff_results_TypeC_YO26-0815-2_ZM99_339.zip"
LEDGER = "YO26-0815-2_ZM99_339.xlsx"


_SUMMARY_ROWS = [
    ("削除図形 総数", 1), ("追加図形 総数", 2), ("変更（追加+削除）図形 総数", 3),
    ("変更なし図形 総数", 4), ("アップロード図面 図形総数", 5), ("図形変更率 [%]", 0.6),
]


def _ledger_bytes():
    from datetime import datetime
    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    ws.append(list(DIFF_LIST_HEADERS))
    ws.append(["C1", "P1", "RevUp", "T", "S", datetime(2026, 10, 18), None, 1, 2, 3, 4, 10, 0])
    summary = wb.create_sheet("Summary")
    for row in _SUMMARY_ROWS:
        summary.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _zip_bytes(prefix):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr(prefix + LEDGER, _ledger_bytes())
    return buf.getvalue()


def _extract_and_find(zip_bytes):
    with tempfile.TemporaryDirectory() as tmp:
        extract_dir = os.path.join(tmp, zip_extract_subdir(ZIP_NAME))
        zipfile.ZipFile(io.BytesIO(zip_bytes)).extractall(extract_dir)
        entries, _ = find_ledger_files(extract_dir)
        return entries


def test_zip_extract_subdir_strips_extension_and_path():
    assert zip_extract_subdir(ZIP_NAME) == ZIP_NAME[:-4]
    assert zip_extract_subdir("a/b\\c.zip") == "c"


@pytest.mark.parametrize("prefix,expected_package", [
    ("", ZIP_NAME[:-4]),  # フォルダ無し ZIP → ZIP のファイル名
    ("inner_folder/", "inner_folder"),  # フォルダ有り ZIP → 内側のフォルダ名
])
def test_package_name_for_flat_and_wrapped_zip(prefix, expected_package):
    entries = _extract_and_find(_zip_bytes(prefix))
    assert [e.package_name for e in entries] == [expected_package]


def test_diff_type_reaches_all_three_sheets_for_flat_zip():
    entries = _extract_and_find(_zip_bytes(""))
    wb = openpyxl.load_workbook(io.BytesIO(build_master_workbook(entries)))
    master = list(wb["Master"].iter_rows(values_only=True))
    work = list(wb["Work Master"].iter_rows(values_only=True))
    summary = list(wb["Summary"].iter_rows(values_only=True))
    assert master[1][master[0].index("Diff Type")] == "C"
    assert work[1][work[0].index("Diff Type")] == "C"
    assert summary[1][summary[0].index("差分方式")] == "C"
