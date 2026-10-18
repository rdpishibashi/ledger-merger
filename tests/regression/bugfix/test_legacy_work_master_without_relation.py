"""
Bugfix regression: 2026-10-18、Relation 列導入前の15列形式の Work Master を持つ
統合図面管理台帳.xlsx を前回分としてアップロードすると、Work Master が読めず
（None）、Summary ともども今回分のみで作成され、蓄積が黙って失われた。

修正: read_work_master_rows() が15列形式（LEGACY_WORK_MASTER_HEADERS_NO_RELATION）を
受け付け、Relation は同じ台帳の Master から (Child, Parent) で補い、
Unchanged Offset Entities は 0 とする。

保証したいこと:
- 15列の Work Master が読め、Relation が Master から補われ、オフセット列が 0 になる
- 前回の Work Master 行が今回分と累積され、Summary にも反映される

実行:
    cd Ledger-merger
    python -m pytest tests/regression/bugfix/test_legacy_work_master_without_relation.py -v
"""
import io
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", ".."))

import openpyxl
from openpyxl import Workbook

from utils.master_ledger_builder import (
    LEGACY_MASTER_HEADERS,
    LEGACY_WORK_MASTER_HEADERS_NO_RELATION,
    WORK_MASTER_HEADERS,
    build_master_workbook,
    read_work_master_rows,
)

_DATE = datetime(2026, 8, 29, 14, 0, 0)


def _legacy_workbook_bytes():
    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    ws.append(LEGACY_MASTER_HEADERS)
    ws.append(["C1", "P1", "RevUp", "T", "S", "A", 1, 2, 3, 4, 10, None, _DATE])
    wm = wb.create_sheet("Work Master")
    wm.append(LEGACY_WORK_MASTER_HEADERS_NO_RELATION)
    # Side が数値で保存されている旧台帳の実態も再現する
    wm.append(["ME24-1001-0", "ZC00", 405, "C1", "P1", "T", "S", "A", 1, 2, 3, 4, 10, None, _DATE])
    wm.append(["ME24-1001-0", "ZC00", 405, "C9", "none", "T", "S", "A", 0, 0, 0, 0, 0, None, _DATE])
    wb.create_sheet("Summary").append(["指番"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_15col_work_master_is_read_with_relation_from_master():
    rows = read_work_master_rows(_legacy_workbook_bytes())
    assert rows is not None and len(rows) == 2
    r = rows[("ME24-1001-0", "ZC00", 405, "C1", "P1")]
    assert len(r) == len(WORK_MASTER_HEADERS)
    assert r[WORK_MASTER_HEADERS.index("Relation")] == "RevUp"
    assert r[WORK_MASTER_HEADERS.index("Unchanged Offset Entities")] == 0
    # Master に無い組は Relation 空欄
    r2 = rows[("ME24-1001-0", "ZC00", 405, "C9", "none")]
    assert r2[WORK_MASTER_HEADERS.index("Relation")] is None


def test_previous_work_master_is_accumulated_not_lost():
    from utils.master_ledger_builder import read_master_rows
    b = _legacy_workbook_bytes()
    out = build_master_workbook(
        [], previous_master_rows=read_master_rows(b), previous_work_master_rows=read_work_master_rows(b),
    )
    wb = openpyxl.load_workbook(io.BytesIO(out))
    assert wb["Work Master"].max_row - 1 == 2
    assert wb["Summary"].max_row - 1 >= 1
