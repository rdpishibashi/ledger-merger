"""不具合再発防止（Bugfix Regression）: Work Masterのキー型ドリフトによる重複行。

不具合ID: 2026-09-01「Work Master シートに Sashiban, Module, Side, Child, Parent が
同じデータが複数ある」（ユーザー報告）

以前どう壊れていたか（再現条件）:
    前回ダウンロードした統合図面管理台帳.xlsx の Work Master シートで、サイドが
    **数値として保存されている**（実データで確認: 統合図面管理台帳_old.xlsx の
    237行すべてが Side=405〈int〉）状態で再アップロードすると、
    read_work_master_rows() が作るキーは ('ME24-1001-0','ZC00',405,...)、今回分の
    キーは ('ME24-1001-0','ZC00','405',...) となり、_merge_by_recorded_date() が
    同一行として認識できずに前回分・今回分の両方が残った
    （実データ: 237行 → 474行、237件が重複）。

    このとき行の**値**だけを str に正規化していたため（当初の実装）、出力上は
    「キー列がまったく同じに見える重複行」になり、原因が分かりにくかった。
    Summaryは2026-09からWork Masterを集計する方式のため、集計値も約2倍に膨れた
    （ユーザー報告の統合図面管理台帳_new.xlsx で確認）。

修正後に保証したいこと:
    - キーは正規化後の行から作り直し、マージ前に両側（前回分・今回分）へ適用する。
      → 型が違っても同一キーとして1行にまとまる。
    - 正規化で同一キーに収束した複数行は "Recorded Date" が最新のものを採用する。

同種の型ドリフトはソートキーでも過去に2度クラッシュを起こしている
（_sort_str() のdocstring、tests/regression/bugfix/test_blank_cell_sort_crash.py 参照）。
値・ソート・**キー**の3箇所すべてで正規化が要る、というのがこの不具合の教訓。
"""

import io
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", ".."))

import openpyxl

from utils.ledger_finder import LedgerEntry
from utils.master_ledger_builder import (
    SUMMARY_HEADERS,
    SUMMARY_SHEET_NAME,
    WORK_MASTER_HEADERS,
    WORK_MASTER_SHEET_NAME,
    build_master_workbook,
)

_PACKAGE_NAME = "dxf_diff_results_TypeA_ME24-1001-0_ZC00_405"


def _row(child, parent, recorded_date, deleted=1, added=2, diff=3, unchanged=4, total=5):
    return (child, parent, "RevUp", "T", "S", recorded_date, None, deleted, added, diff, unchanged, total)


def _previous_rows_with_numeric_side(recorded_date, deleted):
    """サイドが数値（int）で保存された前回のWork Master行を再現する
    （read_work_master_rows() がセルの生値からキーを作るのと同じ形）。"""
    row = (
        "ME24-1001-0", "ZC00", 405, "C1", "P1", "RevUp", "T", "S", "A",
        deleted, 2, 3, 4, 5, None, recorded_date,
    )
    return {("ME24-1001-0", "ZC00", 405, "C1", "P1"): row}


def test_numeric_side_in_previous_ledger_does_not_produce_duplicate_rows():
    """前回分のSideがintでも、今回分（str）と同一キーとして1行にマージされる。"""
    new_entry = LedgerEntry(
        package_name=_PACKAGE_NAME, source_path="new.xlsx",
        diff_list_rows=[_row("C1", "P1", datetime(2026, 9, 1), deleted=11)],
        summary_values={},
    )

    wb_bytes = build_master_workbook(
        [new_entry],
        previous_work_master_rows=_previous_rows_with_numeric_side(datetime(2026, 8, 29), deleted=99),
    )
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    ws = wb[WORK_MASTER_SHEET_NAME]

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    keys = [tuple(row[:5]) for row in rows]
    assert len(rows) == 1, f"重複行が出ている: {rows}"
    assert len(keys) == len(set(keys))
    assert keys[0] == ("ME24-1001-0", "ZC00", "405", "C1", "P1")  # Sideは文字列に健全化
    # Recorded Date が新しい今回分の値が採用される
    assert rows[0][WORK_MASTER_HEADERS.index("Deleted Entities")] == 11


def test_numeric_side_only_previous_rows_are_kept_and_healed():
    """今回分に対応する行が無い前回分も、キー・値ともに正規化されて1行だけ残る。"""
    wb_bytes = build_master_workbook(
        [],
        previous_work_master_rows=_previous_rows_with_numeric_side(datetime(2026, 8, 29), deleted=99),
    )
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    rows = list(wb[WORK_MASTER_SHEET_NAME].iter_rows(min_row=2, values_only=True))

    assert len(rows) == 1
    assert rows[0][:5] == ("ME24-1001-0", "ZC00", "405", "C1", "P1")


def test_summary_totals_are_not_doubled_by_key_type_drift():
    """Summary（2026-09からWork Master由来の集計）の値が重複行で倍にならない。"""
    new_entry = LedgerEntry(
        package_name=_PACKAGE_NAME, source_path="new.xlsx",
        diff_list_rows=[_row("C1", "P1", datetime(2026, 9, 1), deleted=10, added=20, total=100)],
        summary_values={},
    )

    wb_bytes = build_master_workbook(
        [new_entry],
        previous_work_master_rows=_previous_rows_with_numeric_side(datetime(2026, 8, 29), deleted=10),
    )
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    rows = list(wb[SUMMARY_SHEET_NAME].iter_rows(min_row=2, values_only=True))

    assert len(rows) == 1
    row = dict(zip(SUMMARY_HEADERS, rows[0]))
    assert row["変更図面総数"] == 1  # 重複していれば2になる
    assert row["削除図形総数"] == 10  # 同上（20にならない）
    assert row["図形総数"] == 100
