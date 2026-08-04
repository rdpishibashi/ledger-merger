"""仕様確認（Spec Regression）: DXF-diff-manager の2026-08仕様変更への追随。

対応する受入条件（2026-08-02 のユーザー依頼）:
    1. DXF-diff-manager が2026-08以降ZIPダウンロードファイル名末尾のリビジョン番号を
       省略するようになったため、指番/モジュール/サイドの逆算・グルーピングが
       リビジョン無しでも機能する。リビジョン無しのグループはSummaryシートに
       レビジョン別の列を出さず、TOTAL列のみとする。
    2. Master・Work Masterの蓄積マージは、同一キーが前回・今回の両方にある場合、
       "Recorded Date" が新しい方を採用する（今回データによる無条件上書きではない）。
       今回アップロード内で複数フォルダに同一ペアが記録されている場合も同様。
    3. 統合図面管理台帳SummaryのRelation='完全新規図面'の行は「差分ペア総数」から
       除外され、「完全新規図面数」として別集計される（tests/unit/
       test_master_ledger_builder.py::test_compute_summary_rows_excludes_brand_new_from_pair_count
       で詳細に検証済み。ここでは対象外）。
"""

import io
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", ".."))

import openpyxl

from utils.group_summary_builder import (
    build_group_workbook,
    group_entries,
    parse_group_and_revision,
    parse_sashiban_module_side,
)
from utils.ledger_finder import LedgerEntry
from utils.master_ledger_builder import (
    build_master_workbook,
    extract_unique_child_parent_rows,
    extract_unique_work_master_rows,
    read_master_rows,
    read_work_master_rows,
)

# リビジョン省略形（2026-08以降のDXF-diff-manager出力フォルダ名）
_MATCHING_PACKAGE_NAME = "dxf_diff_results_TypeA_ME24-1001-0_ZC00_405"


def _row(child, parent, relation, recorded_date, deleted=1, added=2, diff=3, unchanged=4, total=5):
    return (child, parent, relation, "T", "S", recorded_date, None, deleted, added, diff, unchanged, total)


def test_revision_omitted_package_name_resolves_sashiban_module_side():
    """リビジョン省略形（"..._ZC00_405"）でも指番/モジュール/サイドを逆算できる
    （従来は末尾の "_数字" が必須で (None, None, None) になっていた）。"""
    assert parse_sashiban_module_side(_MATCHING_PACKAGE_NAME) == ("ME24-1001-0", "ZC00", "405")
    assert parse_group_and_revision(_MATCHING_PACKAGE_NAME) == ("ME24-1001-0_ZC00_405", None)


def test_revision_omitted_group_summary_has_total_column_only():
    """グループ内の全エントリがリビジョン省略形の場合、Summaryシートはレビジョン別の
    列を出さず、TOTAL列のみとする。"""
    entry = LedgerEntry(
        package_name=_MATCHING_PACKAGE_NAME, source_path="a.xlsx",
        diff_list_rows=[
            _row("C1", "P1", "RevUp", datetime(2026, 8, 1), deleted=10, added=20, diff=30, unchanged=40, total=100),
        ],
        summary_values={
            "削除図形数 合計": 10, "追加図形数 合計": 20, "差分図形数 合計": 30,
            "変更なし図形数 合計": 40, "総図形数 合計": 100, "図形変更率 [%]": 0.3,
            "入力図面総数": 5, "差分抽出ペア数": 1, "流用率 [%]": 0.2,
            "完全新規図面数": 0, "新規作成率 [%]": 0.0,
        },
    )

    groups = group_entries([entry])
    assert list(groups.keys()) == ["ME24-1001-0_ZC00_405"]
    revision_entries = groups["ME24-1001-0_ZC00_405"]
    assert revision_entries == [(None, entry)]

    wb_bytes = build_group_workbook("ME24-1001-0_ZC00_405", revision_entries)
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    summary_ws = wb["Summary"]

    header = [c.value for c in summary_ws[1]]
    assert header == [None, None, "TOTAL"]  # レビジョン列は無い

    rows_by_label = {row[1]: row[2] for row in summary_ws.iter_rows(min_row=2, values_only=True)}
    assert rows_by_label["完全新規図面数"] == 0
    assert rows_by_label["流用率 [%]"] == 0.2


def test_master_accumulation_merge_keeps_newer_previous_row():
    """前回台帳の行の方が今回アップロード分より Recorded Date が新しい場合、
    今回データによる無条件上書きはせず、前回（より新しい）の行を保持する
    （2026-08、それまでの「今回データが常に前回を上書き」から変更。同一
    Child-Parentペアが複数のDXF-diff-manager出力フォルダに異なる実行時刻で
    記録される実データケースを確認したための修正）。"""
    newer_entry = LedgerEntry(
        package_name=_MATCHING_PACKAGE_NAME, source_path="newer.xlsx",
        diff_list_rows=[_row("C1", "P1", "RevUp", datetime(2026, 8, 1, 22, 0), deleted=999)],
        summary_values={},
    )
    first_bytes = build_master_workbook([newer_entry])
    previous_master_rows = read_master_rows(first_bytes)
    previous_work_master_rows = read_work_master_rows(first_bytes)

    older_entry = LedgerEntry(
        package_name=_MATCHING_PACKAGE_NAME, source_path="older.xlsx",
        diff_list_rows=[_row("C1", "P1", "RevUp", datetime(2026, 8, 1, 17, 0), deleted=1)],
        summary_values={},
    )
    merged_bytes = build_master_workbook(
        [older_entry],
        previous_master_rows=previous_master_rows,
        previous_work_master_rows=previous_work_master_rows,
    )
    wb = openpyxl.load_workbook(io.BytesIO(merged_bytes))

    master_row = next(r for r in wb["Master"].iter_rows(min_row=2, values_only=True) if r[0] == "C1")
    assert master_row[6] == 999  # 前回（より新しいRecorded Date）の値が保持される

    wm_row = next(r for r in wb["Work Master"].iter_rows(min_row=2, values_only=True) if r[3] == "C1")
    assert wm_row[8] == 999  # Work Masterも同様の規則


def test_master_accumulation_merge_overwrites_with_newer_incoming_row():
    """今回アップロード分の方が Recorded Date が新しい場合は、従来どおり今回データで
    上書きする（「新しい方が勝つ」という規則自体は変わっていないことの確認）。"""
    older_entry = LedgerEntry(
        package_name=_MATCHING_PACKAGE_NAME, source_path="older.xlsx",
        diff_list_rows=[_row("C1", "P1", "RevUp", datetime(2026, 8, 1, 17, 0), deleted=1)],
        summary_values={},
    )
    first_bytes = build_master_workbook([older_entry])
    previous_master_rows = read_master_rows(first_bytes)

    newer_entry = LedgerEntry(
        package_name=_MATCHING_PACKAGE_NAME, source_path="newer.xlsx",
        diff_list_rows=[_row("C1", "P1", "RevUp", datetime(2026, 8, 1, 22, 0), deleted=999)],
        summary_values={},
    )
    merged_bytes = build_master_workbook([newer_entry], previous_master_rows=previous_master_rows)
    wb = openpyxl.load_workbook(io.BytesIO(merged_bytes))

    master_row = next(r for r in wb["Master"].iter_rows(min_row=2, values_only=True) if r[0] == "C1")
    assert master_row[6] == 999


def test_within_run_duplicate_pair_picks_latest_recorded_date():
    """今回アップロードした複数フォルダに同一Child-Parentペアが異なる実行時刻で
    記録されている場合（2026-08に実データで確認：同一図番ペアがZC00・ZM00の
    両フォルダに記録され、Recorded Dateが異なっていた）、Masterシート
    （モジュール/サイドを持たない）ではRecorded Dateが最も新しい行が採用される
    （登場順ではなく日時で判定されることを確認する）。

    Work Masterシートはモジュール/サイドをキーに含む（2026-08、Sashiban と Child
    の間に Module・Side 列を追加した際の仕様）ため、この例のように異なるモジュール
    （ZC00・ZM00）に記録された同一Child-Parentは別行として両方残る——このケースは
    重複ではなく、それぞれ別モジュールの実データを表しているため。"""
    entry_appearing_first_but_older = LedgerEntry(
        package_name="dxf_diff_results_TypeA_ME24-1001-0_ZC00_405", source_path="a.xlsx",
        diff_list_rows=[_row("C1", "P1", "RevUp", datetime(2026, 8, 1, 17, 0), deleted=1)],
        summary_values={},
    )
    entry_appearing_second_but_newer = LedgerEntry(
        package_name="dxf_diff_results_TypeA_ME24-1001-0_ZM00_405", source_path="b.xlsx",
        diff_list_rows=[_row("C1", "P1", "RevUp", datetime(2026, 8, 1, 22, 0), deleted=999)],
        summary_values={},
    )

    unique = extract_unique_child_parent_rows(
        [entry_appearing_first_but_older, entry_appearing_second_but_newer]
    )
    assert unique[("C1", "P1")][6] == 999

    unique_wm = extract_unique_work_master_rows(
        [entry_appearing_first_but_older, entry_appearing_second_but_newer]
    )
    assert unique_wm[("ME24-1001-0", "ZC00", "405", "C1", "P1")][8] == 1
    assert unique_wm[("ME24-1001-0", "ZM00", "405", "C1", "P1")][8] == 999
