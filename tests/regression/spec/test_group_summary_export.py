"""仕様確認（Spec Regression）: 指番_モジュール_サイド単位のレビジョン横断集計Excel
（"{指番}_{モジュール}_{サイド}_all.xlsx"）の生成。

対応する受入条件（2026-07-28 のユーザー依頼、添付レイアウト参照ファイル
ME24-1001-0_ZC00_405_all.xlsx との突き合わせで確認）:
    - Summaryシート: TOTAL列 + レビジョン列（"01","02",...）。カウント系5項目は
      レビジョンごとの値をそのまま横に並べ、TOTAL列は単純合計。図形変更率[%]の
      み、TOTAL列は単純合計ではなく TOTAL(分子)/TOTAL(分母)で再計算する。
    - Masterシート: Diff Package列・9項目合計列は含めない。同一Childが複数
      レビジョンにまたがる場合、Deleted/Added/Diff/Unchanged/Total Entitiesの
      5列は単純合計する。
    - 同一出力フォルダに複数の有効な台帳がある場合（実データで確認済みのケース）、
      Recorded Dateが最新のものだけを採用する。

このテストは、ユーザー提供の実際の参照ファイル（ME24-1001-0_ZC00_405_all.xlsx、
2026-07-28受領）のSummaryシートの値と完全一致することを検証する
（tests/fixtures/dxf_diff_manager_output/dxf_diff_results_TypeA_ME24-1001-0_ZC00_405_01〜04
が、その参照ファイルの生成元となった実データそのもの）。

2026-09、ユーザー要求によりSummaryシートの「図面統計」セクション（アップロード
図面総数・差分抽出ペア数・完全新規図面数・流用率[%]・新規作成率[%]）を削除した。
参照ファイル自体にはこのセクションが含まれるが、EXPECTED_SUMMARY_ROWSからは
削除後の現仕様に合わせて対応する行を除いている（「エンティティ統計」6行のみ検証）。
"""

import io
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", ".."))

import openpyxl
import pytest

from utils.group_summary_builder import build_group_workbooks
from utils.ledger_finder import find_ledger_files

REAL_DATA_ROOT = os.path.join(
    os.path.dirname(__file__), "..", "..", "fixtures", "dxf_diff_manager_output"
)

# ユーザー提供の参照ファイル ME24-1001-0_ZC00_405_all.xlsx（2026-07-28）の
# Summaryシートの値をそのまま書き写したもの。
EXPECTED_SUMMARY_ROWS = [
    (None, None, "TOTAL", '"01"', '"02"', '"03"', '"04"'),
    ("エンティティ統計", "削除図形 総数", 518, 222, 233, 10, 53),
    (None, "追加図形 総数", 11652, 2867, 3072, 2835, 2878),
    (None, "変更（追加+削除）図形 総数", 902, 272, 488, 28, 114),
    (None, "変更なし図形 総数", 67705, 31813, 4216, 20708, 10968),
    (None, "アップロード図面 図形総数", 79875, 34902, 7521, 23553, 13899),
    (None, "図形変更率 [%]", 0.01129264475743349, 0.007793249670505988, 0.0648849886983114,
     0.0011888082197596909, 0.008202028922944096),
]


def test_group_summary_matches_user_provided_reference_file():
    entries, _missing = find_ledger_files(REAL_DATA_ROOT)
    files = build_group_workbooks(entries)

    wb = openpyxl.load_workbook(io.BytesIO(files["ME24-1001-0_ZC00_405_all.xlsx"]), data_only=True)
    ws = wb["Summary"]

    actual_rows = [tuple(c.value for c in row) for row in ws.iter_rows(min_row=1, max_row=len(EXPECTED_SUMMARY_ROWS))]

    for actual, expected in zip(actual_rows, EXPECTED_SUMMARY_ROWS):
        assert len(actual) == len(expected)
        for a, e in zip(actual, expected):
            if isinstance(e, float):
                assert a == pytest.approx(e), f"{expected[1]}: {a} != {e}"
            else:
                assert a == e, f"{expected[1]}: {a} != {e}"


def test_group_summary_diff_list_excludes_diff_package_and_total_columns():
    """"Diff Package" 列・合計欄（9項目）は Master シートに含まれない。"""
    entries, _missing = find_ledger_files(REAL_DATA_ROOT)
    files = build_group_workbooks(entries)

    wb = openpyxl.load_workbook(io.BytesIO(files["ME24-1001-0_ZC00_405_all.xlsx"]))
    header = [c.value for c in wb["Master"][1]]

    assert "Diff Package" not in header
    assert len(header) == 12
