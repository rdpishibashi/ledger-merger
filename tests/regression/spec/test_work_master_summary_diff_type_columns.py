"""仕様確認（Spec Regression）: Master・Work Master・Summary シートのデータ形式変更。

対応する受入条件（2026-08-05 のユーザー依頼）:
    Work Master:
        - "Note" を "Total Entities" の後ろへ移動。
        - "Recorded Date" を最後尾へ移動。
        - "Diff Type" を "Deleted Entities" の直前に追加。
    Summary:
        - "差分方式" を "指番" の直後に追加。
    Master（Work Masterと同じ並び替えパターンを適用。フィールド構成自体は
    Work Masterと異なる——Sashiban/Module/Sideは追加せず、Relationはそのまま
    保持する）:
        - "Note" を "Total Entities" の後ろへ移動。
        - "Recorded Date" を最後尾へ移動。
        - "Diff Type" を "Deleted Entities" の直前に追加。
        - ソート順に "Diff Type" を追加（Child → Diff Type の順。Child が全体を
          通じて昇順になることを優先し、Diff Type は同一Childのタイブレークに
          使う。2026-08、当初は逆順〈Diff Type優先〉だったが、ユーザーから
          「Childでソートされていない」と指摘され優先順位を修正）。
        - "Diff Type" フィールドを中央揃いにする（Work Masterも同様）。
    "Diff Type"/"差分方式" は Diff Package（出力フォルダ名。
    "dxf_diff_results_Type{A|B|C}_{指番}(_{モジュール}_{サイド}_*)"）の
    "Type" 直後の1文字から逆算する。

境界・例外の意図:
    - 同一指番内で差分方式が異なる場合（実データでは未確認・理論上のケース）、
      Summaryでは (指番, 差分方式) の組ごとに別行へ分ける
      （ユーザー確認済み仕様。値が誤って混ざらないようにするため）。
    - Work Master・Masterの行自体は差分方式をキーに含めない（1エントリ=1差分方式
      のため行内で曖昧さは生じない。Recorded Date が最も新しい行が丸ごと採用される
      既存の蓄積規則にそのまま従う）。Masterはソート順にのみ差分方式を使う。
    - Masterは指番を逆算できないエントリも対象に含む（Work Masterと異なる）ため、
      Diff Typeの逆算もparse_sashiban_module_side()の解決可否に依存しない。
"""

import io
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", ".."))

import openpyxl

from utils.ledger_finder import LedgerEntry
from utils.master_ledger_builder import (
    MASTER_HEADERS,
    SUMMARY_HEADERS,
    WORK_MASTER_HEADERS,
    build_master_workbook,
    compute_summary_rows,
    extract_unique_child_parent_rows,
    extract_unique_work_master_rows,
)

_TYPE_A_PACKAGE = "dxf_diff_results_TypeA_ME24-1001-0_ZC00_405"
_TYPE_B_PACKAGE = "dxf_diff_results_TypeB_ME24-1001-0_ZM00_405"


def _row(child, parent, relation, recorded_date, deleted=1, added=2, diff=3, unchanged=4, total=5, note=None):
    return (child, parent, relation, "T", "S", recorded_date, note, deleted, added, diff, unchanged, total)


def test_work_master_headers_column_order():
    assert WORK_MASTER_HEADERS == (
        "Sashiban", "Module", "Side", "Child", "Parent", "Title", "Subtitle", "Diff Type",
        "Deleted Entities", "Added Entities", "Diff Entities", "Unchanged Entities",
        "Total Entities", "Note", "Recorded Date",
    )


def test_summary_headers_column_order():
    assert SUMMARY_HEADERS == (
        "指番", "差分方式", "削除図形総数", "追加図形総数", "変更図形総数", "図形総数",
        "図形変更率 [%]", "差分ペア総数", "完全新規図面数", "指番図面総数",
        "流用率 [%]", "新規作成率 [%]", "日付",
    )


def test_work_master_row_has_diff_type_and_moved_note_recorded_date():
    entry = LedgerEntry(
        package_name=_TYPE_A_PACKAGE, source_path="a.xlsx",
        diff_list_rows=[_row("C1", "P1", "RevUp", datetime(2026, 8, 5), note="メモ")],
        summary_values={},
    )

    row = extract_unique_work_master_rows([entry])[("ME24-1001-0", "ZC00", "405", "C1", "P1")]

    assert dict(zip(WORK_MASTER_HEADERS, row)) == {
        "Sashiban": "ME24-1001-0", "Module": "ZC00", "Side": "405", "Child": "C1", "Parent": "P1",
        "Title": "T", "Subtitle": "S", "Diff Type": "A",
        "Deleted Entities": 1, "Added Entities": 2, "Diff Entities": 3, "Unchanged Entities": 4,
        "Total Entities": 5, "Note": "メモ", "Recorded Date": datetime(2026, 8, 5),
    }


def test_summary_row_has_diff_type_after_sashiban():
    entry = LedgerEntry(
        package_name=_TYPE_A_PACKAGE, source_path="a.xlsx",
        diff_list_rows=[_row("C1", "P1", "RevUp", datetime(2026, 8, 5))],
        summary_values={"入力図面総数": 1, "差分抽出ペア数": 1},
    )

    rows = compute_summary_rows([entry], run_timestamp=datetime(2026, 8, 5))

    assert len(rows) == 1
    assert dict(zip(SUMMARY_HEADERS, rows[0]))["指番"] == "ME24-1001-0"
    assert dict(zip(SUMMARY_HEADERS, rows[0]))["差分方式"] == "A"


def test_summary_splits_into_separate_rows_when_diff_type_differs_within_sashiban():
    """同一指番でも差分方式が異なれば、Summaryは (指番,差分方式) ごとに別行になる
    （ユーザー確認済み仕様。実データでは未確認の理論上のケース）。"""
    entry_type_a = LedgerEntry(
        package_name=_TYPE_A_PACKAGE, source_path="a.xlsx",
        diff_list_rows=[_row("C1", "P1", "RevUp", datetime(2026, 8, 5), deleted=10, added=20, total=100)],
        summary_values={"入力図面総数": 4, "差分抽出ペア数": 1},
    )
    entry_type_b = LedgerEntry(
        package_name=_TYPE_B_PACKAGE, source_path="b.xlsx",
        diff_list_rows=[_row("C2", "P2", "RevUp", datetime(2026, 8, 5), deleted=1, added=2, total=10)],
        summary_values={"入力図面総数": 2, "差分抽出ペア数": 1},
    )

    rows = compute_summary_rows([entry_type_a, entry_type_b], run_timestamp=datetime(2026, 8, 5))

    assert len(rows) == 2
    by_type = {row[1]: row for row in rows}
    assert set(by_type.keys()) == {"A", "B"}
    assert by_type["A"][0] == by_type["B"][0] == "ME24-1001-0"  # 指番は同じ
    # 各行の集計値は自分のグループ分のみ（混ざらない）
    assert by_type["A"][2] == 10  # 削除図形総数
    assert by_type["B"][2] == 1
    assert by_type["A"][9] == 4  # 指番図面総数
    assert by_type["B"][9] == 2


def test_master_headers_column_order():
    assert MASTER_HEADERS == (
        "Child", "Parent", "Relation", "Title", "Subtitle", "Diff Type",
        "Deleted Entities", "Added Entities", "Diff Entities", "Unchanged Entities",
        "Total Entities", "Note", "Recorded Date",
    )


def test_master_row_has_diff_type_relation_and_moved_note_recorded_date():
    """Master は Work Master と異なり Sashiban/Module/Side を持たず、Relation を
    そのまま保持する（フィールド構成が異なる、というユーザー指摘どおりの仕様）。"""
    entry = LedgerEntry(
        package_name=_TYPE_A_PACKAGE, source_path="a.xlsx",
        diff_list_rows=[_row("C1", "P1", "RevUp", datetime(2026, 8, 5), note="メモ")],
        summary_values={},
    )

    row = extract_unique_child_parent_rows([entry])[("C1", "P1")]

    assert dict(zip(MASTER_HEADERS, row)) == {
        "Child": "C1", "Parent": "P1", "Relation": "RevUp", "Title": "T", "Subtitle": "S",
        "Diff Type": "A", "Deleted Entities": 1, "Added Entities": 2, "Diff Entities": 3,
        "Unchanged Entities": 4, "Total Entities": 5, "Note": "メモ", "Recorded Date": datetime(2026, 8, 5),
    }


def test_master_diff_type_resolved_even_when_sashiban_unresolvable():
    """Master は Work Master と異なり指番を逆算できないエントリも含むため、
    Diff Type の逆算も指番解決の可否に依存せず独立して行われる。"""
    entry = LedgerEntry(
        package_name=_TYPE_A_PACKAGE, source_path="not_a_ledger_filename.xlsx",
        diff_list_rows=[_row("C1", "P1", "RevUp", datetime(2026, 8, 5))],
        summary_values={},
    )

    row = extract_unique_child_parent_rows([entry])[("C1", "P1")]

    assert dict(zip(MASTER_HEADERS, row))["Diff Type"] == "A"


def test_master_sheet_sorted_by_child_then_diff_type():
    """Masterのソート順は Child が最優先（全体を通じてChild昇順になる）で、
    Diff Type は同一Childが複数の差分方式に跨る場合のタイブレークとしてのみ働く
    （2026-08、ユーザー確認済み仕様。当初は Diff Type を優先していたが、それだと
    Diff Type混在時にChild列が全体としては昇順に見えなくなるとユーザーから指摘され、
    優先順位を逆転した）。差分方式を逆算できないエントリは、その中でのタイブレーク
    上は末尾に回る（が、Childの並び自体には影響しない）。"""
    entry_type_b_c2 = LedgerEntry(
        package_name=_TYPE_B_PACKAGE, source_path="b.xlsx",
        diff_list_rows=[_row("C2", "P2", "RevUp", datetime(2026, 8, 5))],
        summary_values={},
    )
    entry_type_a_c3 = LedgerEntry(
        package_name=_TYPE_A_PACKAGE, source_path="a.xlsx",
        diff_list_rows=[_row("C3", "P3", "RevUp", datetime(2026, 8, 5))],
        summary_values={},
    )
    entry_type_a_c1 = LedgerEntry(
        package_name=_TYPE_A_PACKAGE, source_path="a2.xlsx",
        diff_list_rows=[_row("C1", "P1", "RevUp", datetime(2026, 8, 5))],
        summary_values={},
    )
    entry_type_b_c1 = LedgerEntry(
        package_name=_TYPE_B_PACKAGE, source_path="b2.xlsx",
        diff_list_rows=[_row("C1", "P5", "RevUp", datetime(2026, 8, 5))],
        summary_values={},
    )
    entry_unresolved = LedgerEntry(
        package_name="some_manually_named_folder", source_path="c.xlsx",
        diff_list_rows=[_row("C0", "P0", "RevUp", datetime(2026, 8, 5))],
        summary_values={},
    )

    wb_bytes = build_master_workbook(
        [entry_type_b_c2, entry_type_a_c3, entry_type_a_c1, entry_type_b_c1, entry_unresolved],
    )
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    ws = wb["Master"]

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    # Child が全体を通じて昇順（C0 < C1 < C1 < C2 < C3）。C2(Type B) が C3(Type A) より
    # 前に来ることが、Diff Type優先ではなくChild優先であることの決め手。
    # 同一Child "C1" 内では Diff Type "A" が "B" より先。
    assert [(r[0], r[5]) for r in rows] == [
        ("C0", None), ("C1", "A"), ("C1", "B"), ("C2", "B"), ("C3", "A"),
    ]


def test_master_and_work_master_diff_type_cells_are_centered():
    """Diff Type セルは Master・Work Master いずれも中央揃いにする（2026-08、
    ユーザー要望）。"""
    entry = LedgerEntry(
        package_name=_TYPE_A_PACKAGE, source_path="a.xlsx",
        diff_list_rows=[_row("C1", "P1", "RevUp", datetime(2026, 8, 5))],
        summary_values={},
    )

    wb_bytes = build_master_workbook([entry])
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))

    master_diff_type_col = MASTER_HEADERS.index("Diff Type") + 1
    master_row = next(
        r for r in wb["Master"].iter_rows(min_row=2) if r[0].value == "C1"
    )
    assert master_row[master_diff_type_col - 1].alignment.horizontal == "center"

    wm_diff_type_col = WORK_MASTER_HEADERS.index("Diff Type") + 1
    wm_row = next(
        r for r in wb["Work Master"].iter_rows(min_row=2) if r[3].value == "C1"
    )
    assert wm_row[wm_diff_type_col - 1].alignment.horizontal == "center"
