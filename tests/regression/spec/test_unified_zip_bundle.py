"""仕様確認（Spec Regression）: 統合実行結果を単一ZIP「統合図面台帳.zip」にまとめる。

対応する受入条件（2026-07-28 のユーザー依頼）:
    - ダウンロードボタンは1つ（"統合台帳をダウンロード"）。
    - 出力ZIPのファイル名は固定で "統合図面台帳.zip"。
    - ZIPの中身は次の3種類:
        - "指番_モジュール_サイド別集計/" フォルダ（中のファイル名は変更なし）
        - "図形変更量詳細.xlsx"（旧「統合_図面親子管理台帳_*.xlsx」のリネーム）
        - "統合図面管理台帳.xlsx"（Master: Child-Parentペア蓄積、Work Master: 指番ごとの
          Child-Parentペア蓄積、Summary: 指番ごとの実行時点スナップショットの追記ログ。
          2026-07-31追加）

app.py はこのバンドル処理をそのまま実行するだけの薄いView層なので、ここでは
app.py と同じ手順（build_merged_workbook → build_master_workbook →
build_group_workbooks → zipfile へのまとめ方）を直接再現して検証する。
"""

import io
import os
import sys
import zipfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", ".."))

import openpyxl

from utils.group_summary_builder import build_group_workbooks
from utils.ledger_finder import find_ledger_files
from utils.ledger_merger import build_merged_workbook
from utils.master_ledger_builder import build_master_workbook

REAL_DATA_ROOT = os.path.join(
    os.path.dirname(__file__), "..", "..", "fixtures", "dxf_diff_manager_output"
)


def _build_bundle(entries):
    merged_bytes = build_merged_workbook(entries)
    master_bytes = build_master_workbook(entries)
    group_files = build_group_workbooks(entries)

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("図形変更量詳細.xlsx", merged_bytes)
        zf.writestr("統合図面管理台帳.xlsx", master_bytes)
        for filename, data in sorted(group_files.items()):
            zf.writestr(f"指番_モジュール_サイド別集計/{filename}", data)
    return buffer.getvalue()


def test_bundle_contains_expected_files_with_fixed_names():
    entries, _missing = find_ledger_files(REAL_DATA_ROOT)
    bundle = _build_bundle(entries)

    with zipfile.ZipFile(io.BytesIO(bundle)) as zf:
        names = set(zf.namelist())

    assert "図形変更量詳細.xlsx" in names
    assert "統合図面管理台帳.xlsx" in names
    assert "指番_モジュール_サイド別集計/ME24-1001-0_ZC00_405_all.xlsx" in names
    assert "指番_モジュール_サイド別集計/ME24-1001-0_ZMF1_405_all.xlsx" in names
    assert "指番_モジュール_サイド別集計/ME24-1001-0_ZMB1_405_all.xlsx" in names
    # group_summary_builder.pyが生成するファイル名自体は変更しない（フォルダ名のみ変更）
    assert not any(n.startswith("指番_モジュール_サイド別集計_") for n in names)


def test_master_sheet_in_bundle_is_child_parent_deduped_diff_list_shape():
    from utils.master_ledger_builder import MASTER_HEADERS

    entries, _missing = find_ledger_files(REAL_DATA_ROOT)
    bundle = _build_bundle(entries)

    with zipfile.ZipFile(io.BytesIO(bundle)) as zf:
        master_bytes = zf.read("統合図面管理台帳.xlsx")

    wb = openpyxl.load_workbook(io.BytesIO(master_bytes))
    assert wb.sheetnames == ["Master", "Work Master", "Summary"]
    ws = wb["Master"]
    assert tuple(c.value for c in ws[1]) == MASTER_HEADERS

    child_parent_pairs = [(row[0], row[1]) for row in ws.iter_rows(min_row=2, values_only=True)]
    assert len(child_parent_pairs) == len(set(child_parent_pairs))  # 重複なし
    children = [pair[0] for pair in child_parent_pairs]
    # フィクスチャは全エントリが差分方式"A"で統一のため、並び順はChild昇順と一致する
    # （実際のソートキーは Diff Type→Child。差分方式混在時のソートは
    # test_master_summary_diff_type_columns.py で別途検証）
    assert children == sorted(children)


def test_work_master_sheet_in_bundle_is_sashiban_module_side_child_parent_deduped():
    from utils.master_ledger_builder import WORK_MASTER_HEADERS

    entries, _missing = find_ledger_files(REAL_DATA_ROOT)
    bundle = _build_bundle(entries)

    with zipfile.ZipFile(io.BytesIO(bundle)) as zf:
        master_bytes = zf.read("統合図面管理台帳.xlsx")

    wb = openpyxl.load_workbook(io.BytesIO(master_bytes))
    ws = wb["Work Master"]
    assert tuple(c.value for c in ws[1]) == WORK_MASTER_HEADERS

    keys = [(row[0], row[1], row[2], row[3], row[4]) for row in ws.iter_rows(min_row=2, values_only=True)]
    assert len(keys) == len(set(keys))  # 重複なし（指番,モジュール,サイド,Child,Parent）
    assert keys == sorted(keys)  # 指番→モジュール→サイド→Child昇順
    assert all(sashiban is not None for sashiban, _m, _s, _child, _parent in keys)  # 指番不明エントリは除外済み


def test_summary_sheet_in_bundle_has_one_row_per_sashiban_for_first_run():
    from utils.master_ledger_builder import SUMMARY_HEADERS

    entries, _missing = find_ledger_files(REAL_DATA_ROOT)
    bundle = _build_bundle(entries)

    with zipfile.ZipFile(io.BytesIO(bundle)) as zf:
        master_bytes = zf.read("統合図面管理台帳.xlsx")

    wb = openpyxl.load_workbook(io.BytesIO(master_bytes))
    ws = wb["Summary"]
    assert tuple(c.value for c in ws[1]) == SUMMARY_HEADERS

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    # フィクスチャは単一指番（ME24-1001-0）のため、初回実行では1行のみ
    assert len(rows) == 1
    assert rows[0][0] == "ME24-1001-0"
