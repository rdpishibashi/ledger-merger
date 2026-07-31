"""utils.group_summary_builder のユニットテスト。

実データ（tests/fixtures/dxf_diff_manager_output、tests/unit/test_ledger_merger.py の
モジュールdocstring参照）と、集計ロジック（Childの重複・'n/a'混在・同一フォルダ内の
複数有効台帳）を検証するための合成データの両方を使う。
"""

import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))

from utils.group_summary_builder import (
    aggregate_diff_list_by_child,
    build_group_workbooks,
    group_entries,
    parse_group_and_revision,
    parse_sashiban_module_side,
)
from utils.ledger_finder import LedgerEntry, find_ledger_files

REAL_DATA_ROOT = os.path.join(
    os.path.dirname(__file__), "..", "fixtures", "dxf_diff_manager_output"
)


def _row(child, parent, relation, recorded_date, deleted, added, diff, unchanged, total):
    return (child, parent, relation, "T", "S", recorded_date, None, deleted, added, diff, unchanged, total)


def test_parse_group_and_revision():
    assert parse_group_and_revision("dxf_diff_results_PairA_ME24-1001-0_ZC00_405_04") == (
        "ME24-1001-0_ZC00_405", "04",
    )
    assert parse_group_and_revision("dxf_diff_results_PairC_ME24-1001-0_ZMB1_405_01") == (
        "ME24-1001-0_ZMB1_405", "01",
    )


def test_parse_group_and_revision_accepts_new_type_naming():
    """DXF-diff-manager が2026-07-28以降に自動生成する命名規則（"Type"トークン）にも対応する。"""
    assert parse_group_and_revision("dxf_diff_results_TypeA_ME24-1001-0_ZC00_405_04") == (
        "ME24-1001-0_ZC00_405", "04",
    )


def test_parse_group_and_revision_returns_none_for_unrelated_names():
    assert parse_group_and_revision("some_other_folder") is None
    assert parse_group_and_revision("dxf_diff_results_PairA_no_revision_suffix") is None


def test_parse_sashiban_module_side():
    assert parse_sashiban_module_side("dxf_diff_results_PairA_ME24-1001-0_ZC00_405_04") == (
        "ME24-1001-0", "ZC00", "405",
    )
    assert parse_sashiban_module_side("dxf_diff_results_TypeB_ME24-1001-0_ZMF1_405_01") == (
        "ME24-1001-0", "ZMF1", "405",
    )


def test_parse_sashiban_module_side_accepts_na_module_and_side():
    """DXF-diff-manager でモジュール/サイドが未入力の場合、"na" のまま返す。"""
    assert parse_sashiban_module_side("dxf_diff_results_TypeA_ME24-1001-0_na_na_01") == (
        "ME24-1001-0", "na", "na",
    )


def test_parse_sashiban_module_side_returns_none_tuple_for_unrelated_names():
    assert parse_sashiban_module_side("some_other_folder") == (None, None, None)
    assert parse_sashiban_module_side("dxf_diff_results_PairA_no_revision_suffix") == (None, None, None)


def test_group_entries_picks_latest_recorded_date_when_folder_has_duplicate_ledgers():
    """同一フォルダ（同一package_name）に複数の有効な台帳がある場合、Diff List内の
    最大Recorded Dateが最も新しいものだけを採用する
    （2026-07-28に実データで確認: 図番抽出に失敗した古い実行結果と、成功した新しい
    実行結果が同じフォルダに混在していたケース）。"""
    older = LedgerEntry(
        package_name="dxf_diff_results_PairA_ME00-0000-0_ZZ00_000_01",
        source_path="old.xlsx",
        diff_list_rows=[_row("C1", "P1", "RevUp", datetime(2026, 7, 7), 1, 1, 2, 1, 3)],
        summary_values={},
    )
    newer = LedgerEntry(
        package_name="dxf_diff_results_PairA_ME00-0000-0_ZZ00_000_01",
        source_path="new.xlsx",
        diff_list_rows=[_row("C1", "P1", "RevUp", datetime(2026, 7, 9), 9, 9, 18, 9, 27)],
        summary_values={},
    )

    groups = group_entries([older, newer])
    assert list(groups.keys()) == ["ME00-0000-0_ZZ00_000"]
    revisions = groups["ME00-0000-0_ZZ00_000"]
    assert len(revisions) == 1
    revision, chosen_entry = revisions[0]
    assert revision == "01"
    assert chosen_entry.source_path == "new.xlsx"


def test_aggregate_diff_list_by_child_sums_entity_columns_across_revisions():
    """仕様どおり、同一Childが複数レビジョンにまたがる場合は5列を単純合計する。"""
    entry_rev1 = LedgerEntry(
        package_name="p_01", source_path="r1.xlsx",
        diff_list_rows=[
            _row("C1", "P1", "RevUp", datetime(2026, 7, 1), 1, 2, 3, 4, 5),
            _row("C2", "none", "完全新規図面", datetime(2026, 7, 1), "n/a", 10, "n/a", "n/a", 10),
        ],
        summary_values={},
    )
    entry_rev2 = LedgerEntry(
        package_name="p_02", source_path="r2.xlsx",
        diff_list_rows=[
            _row("C1", "P1", "RevUp", datetime(2026, 7, 2), 10, 20, 30, 40, 50),
            _row("C2", "none", "完全新規図面", datetime(2026, 7, 2), "n/a", 10, "n/a", "n/a", 10),
        ],
        summary_values={},
    )

    aggregated = aggregate_diff_list_by_child([("01", entry_rev1), ("02", entry_rev2)])

    by_child = {row[0]: row for row in aggregated}
    assert by_child["C1"][7:] == (11, 22, 33, 44, 55)
    # C2 は全出現が 'n/a' の列（Deleted/Diff/Unchanged）はそのまま 'n/a'、Added/Totalは合計
    assert by_child["C2"][7:] == ("n/a", 20, "n/a", "n/a", 20)


def test_aggregate_diff_list_by_child_uses_latest_row_for_non_entity_columns():
    """非数値列（Parent/Relation/Title等）は最新のRecorded Dateの行の値を採用する。"""
    entry_rev1 = LedgerEntry(
        package_name="p_01", source_path="r1.xlsx",
        diff_list_rows=[("C1", "P_OLD", "流用", "T_OLD", "S_OLD", datetime(2026, 7, 1), None, 1, 1, 2, 1, 3)],
        summary_values={},
    )
    entry_rev2 = LedgerEntry(
        package_name="p_02", source_path="r2.xlsx",
        diff_list_rows=[("C1", "P_NEW", "RevUp", "T_NEW", "S_NEW", datetime(2026, 7, 2), None, 1, 1, 2, 1, 3)],
        summary_values={},
    )

    aggregated = aggregate_diff_list_by_child([("01", entry_rev1), ("02", entry_rev2)])
    assert len(aggregated) == 1
    row = aggregated[0]
    assert row[1:5] == ("P_NEW", "RevUp", "T_NEW", "S_NEW")
    assert row[5] == datetime(2026, 7, 2)


def test_build_group_workbooks_from_real_fixture():
    """実データフィクスチャからグループ別Excelが生成され、Summary/Diff Listの
    構造が想定通りであることを確認する。"""
    entries, _missing = find_ledger_files(REAL_DATA_ROOT)
    files = build_group_workbooks(entries)

    assert set(files.keys()) == {
        "ME24-1001-0_ZMF1_405_all.xlsx",
        "ME24-1001-0_ZMB1_405_all.xlsx",
        "ME24-1001-0_ZC00_405_all.xlsx",
    }

    import io
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(files["ME24-1001-0_ZMF1_405_all.xlsx"]))
    assert wb.sheetnames == ["Summary", "Diff List"]

    summary_ws = wb["Summary"]
    header = [c.value for c in summary_ws[1]]
    assert header == [None, None, "TOTAL", '"01"']

    diff_ws = wb["Diff List"]
    diff_header = tuple(c.value for c in diff_ws[1])
    from utils.ledger_finder import DIFF_LIST_HEADERS
    assert diff_header == DIFF_LIST_HEADERS
    assert diff_ws.max_row - 1 == 2  # ZMF1_405_01 は2行（Diff Package・合計列は含まない）


def test_diff_list_entity_columns_are_formatted_and_centered():
    """"* Entities" 列はカンマ区切り＋中央揃いで表示する（'n/a' と数値の位置ずれ対策）。
    ヘッダー行も中央揃いにする。"""
    import io

    import openpyxl

    from utils.ledger_finder import DIFF_LIST_HEADERS

    entries, _missing = find_ledger_files(REAL_DATA_ROOT)
    files = build_group_workbooks(entries)

    wb = openpyxl.load_workbook(io.BytesIO(files["ME24-1001-0_ZC00_405_all.xlsx"]))
    diff_ws = wb["Diff List"]

    for cell in diff_ws[1]:
        assert cell.alignment.horizontal == "center"

    entity_cols = [DIFF_LIST_HEADERS.index(label) + 1 for label in
                   ("Deleted Entities", "Added Entities", "Diff Entities", "Unchanged Entities", "Total Entities")]
    for row_idx in range(2, diff_ws.max_row + 1):
        for col in entity_cols:
            cell = diff_ws.cell(row=row_idx, column=col)
            assert cell.number_format == "#,##0"
            assert cell.alignment.horizontal == "center"
