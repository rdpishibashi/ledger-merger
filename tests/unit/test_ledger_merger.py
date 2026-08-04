"""実データ（tests/fixtures/dxf_diff_manager_output、コミット済み）を使った
Ledger-merger の回帰テスト。

このフィクスチャは 2026-07-28 に実際に報告された不具合の調査時に取得した実データの
一部で、DXF-diff-manager の現行の出力形式（Summary シートのラベル文言、
`dxf図面` サブフォルダ付きの出力フォルダ構成）をそのまま反映している。

フィクスチャ構成:
    dxf_diff_manager_output/
    ├── ME24-1001-0_ZM00_405_all.xlsx
    │       ZIP直下（ラッパー）に置かれた、有効な台帳フォーマットだが統合対象外の
    │       集約ファイル。より深い階層の各出力フォルダにもxlsxがあるため、
    │       中間ラッパーとして無視されるべき（ユーザー確認済みの仕様）。
    ├── dxf_diff_results_PairA_ME24-1001-0_ZMF1_405_01/
    │       有効な台帳（2行）+ dxf図面/（xlsxを含まないサブフォルダ）付き。
    ├── dxf_diff_results_PairA_ME24-1001-0_ZMB1_405_01/
    │       有効な台帳（3行）。
    ├── dxf_diff_results_PairA_ME24-1001-0_ZMF2_405_01/
    │       diff_labels.xlsx / unchanged_labels.xlsx のみで有効な台帳が無い
    │       （「台帳が見つからなかったフォルダ」として報告されるべき）。
    └── dxf_diff_results_PairA_ME24-1001-0_ZC00_405_01〜04/
            "ME24-1001-0_ZC00_405"（指番_モジュール_サイド）1グループ・4レビジョン分。
            `_02` のみ有効な台帳が2件（ME24-1001-0_ZC00_405.xlsx と
            ME24-1001-0_na_na.xlsx、後者は図番抽出に失敗した古い実行結果の残骸）。
            utils.group_summary_builder のグルーピング・重複排除・レビジョン横断
            集計の回帰テスト（tests/regression/spec/test_group_summary_export.py）に使う。
"""

import os
import shutil
import sys

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))

from utils.group_summary_builder import parse_sashiban_module_side
from utils.ledger_finder import LedgerEntry, find_ledger_files, reconcile_missing_folders
from utils.ledger_merger import OUTPUT_HEADERS, build_merged_workbook

REAL_DATA_ROOT = os.path.join(
    os.path.dirname(__file__), "..", "fixtures", "dxf_diff_manager_output"
)
SAMPLE_LEDGER_DIR = os.path.join(
    REAL_DATA_ROOT, "dxf_diff_results_PairA_ME24-1001-0_ZMF1_405_01"
)
SAMPLE_LEDGER_FILENAME = "ME24-1001-0_ZMF1_405.xlsx"

MISSING_LEDGER_FOLDER_NAME = "dxf_diff_results_PairA_ME24-1001-0_ZMF2_405_01"


def test_find_ledger_files_count_matches_actual_files():
    """フィクスチャ内の有効な台帳は7件（ZMF1_405_01, ZMB1_405_01, ZC00_405_01〜04。
    ただし ZC00_405_02 のみ2件）。ZM00_405_all.xlsx（ラッパー直下の集約ファイル）と
    ZMF2_405_01（台帳無し）は entries に含まれない。"""
    entries, missing_folders = find_ledger_files(REAL_DATA_ROOT)

    assert len(entries) == 7
    assert {e.package_name for e in entries} == {
        "dxf_diff_results_PairA_ME24-1001-0_ZMF1_405_01",
        "dxf_diff_results_PairA_ME24-1001-0_ZMB1_405_01",
        "dxf_diff_results_PairA_ME24-1001-0_ZC00_405_01",
        "dxf_diff_results_PairA_ME24-1001-0_ZC00_405_02",
        "dxf_diff_results_PairA_ME24-1001-0_ZC00_405_03",
        "dxf_diff_results_PairA_ME24-1001-0_ZC00_405_04",
    }
    assert missing_folders == [MISSING_LEDGER_FOLDER_NAME]


def test_non_ledger_filenames_excluded_and_not_reported_as_missing():
    """diff_labels.xlsx / unchanged_labels.xlsx は台帳候補から除外され、
    台帳が別途存在するフォルダは「台帳が見つからないフォルダ」にも出てこない。"""
    entries, missing_folders = find_ledger_files(REAL_DATA_ROOT)

    ledger_paths = {entry.source_path for entry in entries}
    assert all(os.path.basename(p) not in {"diff_labels.xlsx", "unchanged_labels.xlsx"} for p in ledger_paths)

    found_names = {e.package_name for e in entries}
    assert found_names.isdisjoint(missing_folders)


def test_folder_with_only_non_ledger_files_is_reported_as_missing_by_name_only(tmp_path):
    """diff_labels.xlsx / unchanged_labels.xlsx しか無いフォルダは「台帳が見つからない
    フォルダ」として basename のみで報告され、台帳が存在するフォルダは混在しても無視されない。"""
    ok_dir = tmp_path / "dxf_diff_results_PairC_OK_01"
    ok_dir.mkdir()
    shutil.copy(
        os.path.join(SAMPLE_LEDGER_DIR, SAMPLE_LEDGER_FILENAME),
        ok_dir / SAMPLE_LEDGER_FILENAME,
    )

    missing_dir = tmp_path / "dxf_diff_results_PairC_MISSING_01"
    missing_dir.mkdir()
    shutil.copy(os.path.join(SAMPLE_LEDGER_DIR, "diff_labels.xlsx"), missing_dir / "diff_labels.xlsx")
    shutil.copy(os.path.join(SAMPLE_LEDGER_DIR, "unchanged_labels.xlsx"), missing_dir / "unchanged_labels.xlsx")

    entries, missing_folders = find_ledger_files(str(tmp_path))

    assert len(entries) == 1
    assert entries[0].package_name == "dxf_diff_results_PairC_OK_01"
    assert missing_folders == ["dxf_diff_results_PairC_MISSING_01"]


def test_rows_without_diff_stats_are_excluded(tmp_path):
    """Diff List には「差分抽出ペア数」と同数の行（実際に差分抽出された行）だけが
    含まれる。Total Entities が空欄の行（差分未抽出の図番ペアの関係記録）は除外される。

    現行のコミット済みフィクスチャには空欄行を含む実データが無いため、
    DXF-diff-manager の現行フォーマットに沿った最小限の台帳ファイルをその場で
    合成して検証する（決定的なテストにするため）。"""
    from openpyxl import Workbook

    path = tmp_path / "dxf_diff_results_PairX_ME00-0000-0_01" / "ME00-0000-0_ZZ00_000.xlsx"
    path.parent.mkdir(parents=True)

    wb = Workbook()
    diff_ws = wb.active
    diff_ws.title = "Diff List"
    diff_ws.append([
        "Child", "Parent", "Relation", "Title", "Subtitle", "Recorded Date", "Note",
        "Deleted Entities", "Added Entities", "Diff Entities", "Unchanged Entities",
        "Total Entities",
    ])
    diff_ws.append(["C1", "P1", "RevUp", "T", "S", None, None, 1, 2, 3, 4, 10])
    diff_ws.append(["C2", "P2", "親子", "T", "S", None, None, None, None, None, None, None])
    diff_ws.append(["C3", "P3", "RevUp", "T", "S", None, None, 5, 6, 11, 7, 25])

    summary_ws = wb.create_sheet("Summary")
    for row in [
        ("削除図形 総数", 6), ("追加図形 総数", 8), ("変更（追加+削除）図形 総数", 14),
        ("変更なし図形 総数", 11), ("アップロード図面 図形総数", 35),
        ("図形変更率 [%]", 0.4), ("アップロード図面総数", 3),
        ("差分抽出ペア数", 2), ("流用率 [%]", 0.66),
    ]:
        summary_ws.append(row)

    wb.save(path)

    entries, _missing_folders = find_ledger_files(str(tmp_path))
    assert len(entries) == 1
    entry = entries[0]

    assert len(entry.diff_list_rows) == entry.summary_values["差分抽出ペア数"] == 2
    assert all(row[-1] is not None for row in entry.diff_list_rows)
    assert [row[0] for row in entry.diff_list_rows] == ["C1", "C3"]


def test_filtered_rows_entity_sums_match_summary_exactly():
    """除外後に残った行の Deleted/Added/Diff/Unchanged/Total Entities の合計が、
    Summary シートの対応する合計値と1件単位の差もなく完全に一致することを、
    件数だけでなく数値レベルで検証する。完全新規図面の行は該当列が 'n/a'（文字列）に
    なるため、DXF-diff-manager 自身の集計と同様に数値以外はスキップして合計する。"""
    from utils.ledger_finder import DIFF_LIST_HEADERS

    cols = {h: i for i, h in enumerate(DIFF_LIST_HEADERS)}
    label_by_col = {
        "Deleted Entities": "削除図形数 合計",
        "Added Entities": "追加図形数 合計",
        "Diff Entities": "差分図形数 合計",
        "Unchanged Entities": "変更なし図形数 合計",
        "Total Entities": "総図形数 合計",
    }

    entries, _missing_folders = find_ledger_files(REAL_DATA_ROOT)
    assert entries

    for entry in entries:
        for col_name, summary_label in label_by_col.items():
            computed = sum(
                row[cols[col_name]] for row in entry.diff_list_rows
                if isinstance(row[cols[col_name]], (int, float))
            )
            assert computed == entry.summary_values[summary_label], (
                f"{entry.package_name}: {summary_label} 不一致"
            )


def test_filtered_rows_preserve_original_order_and_values():
    """除外後に残る行は、元の Diff List シートの該当行を順序・値ともに
    そのまま保持している（並び替えやデータ欠落が無いことの往復確認）。"""
    from utils.ledger_finder import _TOTAL_ENTITIES_COL

    entries, _missing_folders = find_ledger_files(REAL_DATA_ROOT)
    for entry in entries:
        wb = openpyxl.load_workbook(entry.source_path, data_only=True, read_only=True)
        try:
            src_rows = list(wb["Diff List"].iter_rows(values_only=True))[1:]
        finally:
            wb.close()
        expected = [row for row in src_rows if row[_TOTAL_ENTITIES_COL] is not None]
        assert list(entry.diff_list_rows) == expected


def test_reconcile_missing_folders_removes_names_found_elsewhere():
    """異なる入力ソース（複数ZIP等）の結果を集約した際、同名フォルダが片方で台帳あり・
    もう片方で台帳なしと判定されても、「台帳が見つからなかったフォルダ」一覧には
    矛盾して現れない（成功した名前を優先する）。重複した欠落名も1つにまとめる。"""
    entries = [LedgerEntry(package_name="A", source_path="x", diff_list_rows=[], summary_values={})]
    missing = ["A", "B", "B", "C"]

    reconciled = reconcile_missing_folders(entries, missing)

    assert reconciled == ["B", "C"]


def test_macosx_mirror_folder_not_reported_as_missing(tmp_path):
    """macOS Finder/ditto で ZIP 化すると同名の __MACOSX/<folder>/._ファイル名 ミラーが
    作られる。これが本物のフォルダと同名の偽フォルダとして誤検出されないことを確認する
    （実際に発生した不具合: 台帳が見つかるフォルダ名がそのまま「見つからなかったフォルダ」
    にも重複して表示されていた）。"""
    real_dir = tmp_path / "dxf_diff_results_PairC_OK_01"
    real_dir.mkdir()
    shutil.copy(
        os.path.join(SAMPLE_LEDGER_DIR, SAMPLE_LEDGER_FILENAME),
        real_dir / SAMPLE_LEDGER_FILENAME,
    )

    mirror_dir = tmp_path / "__MACOSX" / "dxf_diff_results_PairC_OK_01"
    mirror_dir.mkdir(parents=True)
    (mirror_dir / f"._{SAMPLE_LEDGER_FILENAME}").write_bytes(b"\x00\x05\x16\x07")  # AppleDouble ダミー

    entries, missing_folders = find_ledger_files(str(tmp_path))

    assert len(entries) == 1
    assert entries[0].package_name == "dxf_diff_results_PairC_OK_01"
    assert missing_folders == []


def test_merged_workbook_structure():
    entries, _missing_folders = find_ledger_files(REAL_DATA_ROOT)
    merged_bytes = build_merged_workbook(entries)

    wb = openpyxl.load_workbook(__import__("io").BytesIO(merged_bytes))
    ws = wb["Diff List"]

    header = tuple(c.value for c in ws[1])
    assert header == OUTPUT_HEADERS
    assert len(header) == 21
    assert header[:3] == ("Sashiban", "Module", "Side")
    assert header[-1] == "Diff Package"
    # 図面統計系（入力図面総数・差分抽出ペア数・流用率 [%]・変更なし図形数 合計）は
    # 含まれない。表示名も変更されている（差分図形数 合計→変更図形数 合計、
    # 総図形数 合計→図形総数 合計）。
    assert "変更なし図形数 合計" not in header
    assert "入力図面総数" not in header
    assert "差分抽出ペア数" not in header
    assert "流用率 [%]" not in header
    assert "変更図形数 合計" in header
    assert "図形総数 合計" in header

    row_idx = 2
    for entry in entries:
        sashiban, module, side = parse_sashiban_module_side(entry.package_name, entry.source_path)
        for row_in_block in range(len(entry.diff_list_rows)):
            row = ws[row_idx]
            package_cell = row[-1]
            summary_cells = row[15:-1]

            if row_in_block == 0:
                assert package_cell.font.color.rgb == "FF000000"
                assert all(c.value is not None for c in summary_cells)
            else:
                assert package_cell.font.color.rgb == "FFA6A6A6"
                assert all(c.value is None for c in summary_cells)

            assert package_cell.value == entry.package_name
            assert (row[0].value, row[1].value, row[2].value) == (sashiban, module, side)
            row_idx += 1

    assert row_idx - 2 == sum(len(e.diff_list_rows) for e in entries)


def _diff_row(child, parent):
    return (child, parent, "RevUp", "T", "S", None, None, 1, 2, 3, 4, 10)


_MERGED_SUMMARY_VALUES = {
    "削除図形数 合計": 1, "追加図形数 合計": 2, "差分図形数 合計": 3,
    "総図形数 合計": 10, "図形変更率 [%]": 0.3,
}


def test_merged_workbook_rows_sorted_by_sashiban_diff_package_module_side_child():
    """行は Sashiban → Diff Package → Module → Side → Child の昇順に並ぶ
    （2026-08、ユーザー要望）。ブロック内で最初に現れる行（＝ソート後にChildが
    最小の行）に集計値・Diff Package列の黒字フォントが付くことも確認する。"""
    entry_zz_c3_c1 = LedgerEntry(
        package_name="dxf_diff_results_TypeA_ZZ99-0001-0_ZM00_405", source_path="a.xlsx",
        diff_list_rows=[_diff_row("C3", "P3"), _diff_row("C1", "P1")],
        summary_values=_MERGED_SUMMARY_VALUES,
    )
    entry_aa_c2_c1 = LedgerEntry(
        package_name="dxf_diff_results_TypeA_AA10-0001-0_ZM00_405", source_path="b.xlsx",
        diff_list_rows=[_diff_row("C2", "P2"), _diff_row("C1", "P1")],
        summary_values=_MERGED_SUMMARY_VALUES,
    )

    merged_bytes = build_merged_workbook([entry_zz_c3_c1, entry_aa_c2_c1])
    wb = openpyxl.load_workbook(__import__("io").BytesIO(merged_bytes))
    ws = wb["Diff List"]

    child_col = OUTPUT_HEADERS.index("Child") + 1
    sashiban_col = OUTPUT_HEADERS.index("Sashiban") + 1
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert [(r[sashiban_col - 1], r[child_col - 1]) for r in rows] == [
        ("AA10-0001-0", "C1"), ("AA10-0001-0", "C2"),
        ("ZZ99-0001-0", "C1"), ("ZZ99-0001-0", "C3"),
    ]

    # 各ブロックの先頭行（Child最小）のみ集計値あり・Diff Package列が黒字
    summary_col = OUTPUT_HEADERS.index("削除図形数 合計") + 1
    for row_idx, expected_first in zip(range(2, ws.max_row + 1), [True, False, True, False]):
        summary_cell_value = ws.cell(row=row_idx, column=summary_col).value
        package_font = ws.cell(row=row_idx, column=OUTPUT_HEADERS.index("Diff Package") + 1).font.color.rgb
        if expected_first:
            assert summary_cell_value == 1
            assert package_font == "FF000000"
        else:
            assert summary_cell_value is None
            assert package_font == "FFA6A6A6"


def test_merged_workbook_unresolvable_sashiban_sorts_last():
    """Sashiban/Module/Side を逆算できないエントリ（命名規則に一致しない）は、
    ソート時に空欄扱いのため他の解決済みエントリより後ろに並ぶ。"""
    entry_resolved = LedgerEntry(
        package_name="dxf_diff_results_TypeA_AA10-0001-0_ZM00_405", source_path="a.xlsx",
        diff_list_rows=[_diff_row("C1", "P1")],
        summary_values=_MERGED_SUMMARY_VALUES,
    )
    entry_unresolved = LedgerEntry(
        package_name="some_manually_named_folder", source_path="b.xlsx",
        diff_list_rows=[_diff_row("C1", "P1")],
        summary_values=_MERGED_SUMMARY_VALUES,
    )

    merged_bytes = build_merged_workbook([entry_unresolved, entry_resolved])
    wb = openpyxl.load_workbook(__import__("io").BytesIO(merged_bytes))
    ws = wb["Diff List"]

    sashiban_values = [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert sashiban_values == ["AA10-0001-0", None]


def test_entity_and_summary_columns_are_formatted_and_centered():
    """"* Entities" 列・Summary由来5項目（"* 合計" 等）列はカンマ区切り（％項目は
    0.00%）＋中央揃いで表示する。'n/a' と数値の位置がずれないようにするため。
    ヘッダー行も中央揃いにする。"""
    from utils.ledger_merger import ENTITY_COLS, OUTPUT_HEADERS, PERCENT_LABELS, _MERGED_SUMMARY_DISPLAY_LABELS

    entries, _missing_folders = find_ledger_files(REAL_DATA_ROOT)
    merged_bytes = build_merged_workbook(entries)

    wb = openpyxl.load_workbook(__import__("io").BytesIO(merged_bytes))
    ws = wb["Diff List"]

    for cell in ws[1]:
        assert cell.alignment.horizontal == "center"

    summary_start_col = OUTPUT_HEADERS.index(_MERGED_SUMMARY_DISPLAY_LABELS[0]) + 1
    row_idx = 2
    for entry in entries:
        for row_in_block in range(len(entry.diff_list_rows)):
            for col in ENTITY_COLS:
                cell = ws.cell(row=row_idx, column=col)
                assert cell.number_format == "#,##0"
                assert cell.alignment.horizontal == "center"

            if row_in_block == 0:
                for offset, label in enumerate(_MERGED_SUMMARY_DISPLAY_LABELS):
                    cell = ws.cell(row=row_idx, column=summary_start_col + offset)
                    expected_format = "0.00%" if label in PERCENT_LABELS else "#,##0"
                    assert cell.number_format == expected_format
                    assert cell.alignment.horizontal == "center"

            row_idx += 1
